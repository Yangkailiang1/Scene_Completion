"""V2 missing-concern Markdown/Excel and graph exports."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
from .interaction_protocol import baseline
from .interaction_reports import write_diagrams, refs_text
from .semantic_backend import write_json
from .sources import target_sections

LABELS={'miss_rate':'缺失异常漏报率','generation_validity':'生成有效率','adoption_rate':'推荐采纳率',
    'recommendation_false_positive_rate':'推荐误报率（分母GT）'}
GROUP_LABELS={'api_data':'传输数据','common':'通用交互','human':'人类参与者','external_service':'外部服务',
    'external_database':'外部数据库资源服务','internal_database':'内部数据库资源服务',
    'external_llm':'外部LLM调用','external_llm_quality':'LLM质量','service_relation':'内部服务依赖关系',
    'service_dependency':'内部依赖可用性','service_display_interaction':'展示交互型服务',
    'service_query_retrieval':'查询检索型服务','service_resource_mutation':'资源变更型服务',
    'service_analysis_generation':'分析生成型服务','service_release_activation':'发布生效型服务',
    'service_workflow':'工作流','unclassified':'未分类异常'}


def export(root,bundle,accepted=False):
    root=Path(root);root.mkdir(parents=True,exist_ok=True)
    values=bundle['metrics'];defs=baseline(); ranking=bundle['ranking']
    units={u['scenario_id']:u for u in bundle['G']}
    missing=[units[i] for i in values['missing_ids']]
    selected=set(ranking['selected_ids'])
    status='方法流程验收通过' if accepted else '流程结果，尚未完成最终执行与独立审查验收'
    lines=['# 交互关注点缺失报告','',status,'',
        '范围：'+', '.join(bundle['scope']['use_case_ids'])+'；'+bundle['scope']['granularity'],
        'GT暂取客户测试spec；明示与框架推导细粒度原因均计入，并分别标记。',
        '推荐分数为同类历史余弦相似度原始求和，不是正确概率；无历史项使用固定种子随机抽样。','',
        '| 指标 | 分子 | 分母 | 值 |','| --- | ---: | ---: | ---: |']
    for metric in values['overall']['metrics']:
        value='暂无分母' if metric['value'] is None else f"{metric['value']:.2%}"
        lines.append(f"| {LABELS[metric['metric']]} | {metric['numerator']} | {metric['denominator']} | {value} |")
    lines+=['','## 关注点与大类指标','',
        '| 类别 | 关注点说明 | 指标 | 分子/分母 | 值 |','| --- | --- | --- | ---: | ---: |']
    for bucket in values['concerns']+values['groups']+[values['unclassified']]:
        key=bucket['category']
        desc=defs[key]['label'] if key in defs else GROUP_LABELS.get(key,key)
        for m in bucket['metrics']:
            if not m['denominator'] and not m['numerator']:continue
            value='暂无分母' if m['value'] is None else f"{m['value']:.2%}"
            lines.append(f"| {key} | {desc} | {LABELS[m['metric']]} | {m['numerator']}/{m['denominator']} | {value} |")
    lines+=['','## 缺失的交互关注点异常','']
    rows=[]
    for u in missing:
        loc=u['location_id'] or '未映射'
        refs=[r for rs in target_sections(bundle['source_index'],u['use_case_id']).values() for r in rs]
        target=refs_text(refs) if refs else '章节定位待确认'
        lines.append(f"- {u['use_case_id']} / {loc} / {', '.join(u['concern_keys'])}：{u['trigger_condition']}；建议补充于 {target}")
        rows.append({'unit_id':u['scenario_id'],'use_case_id':u['use_case_id'],'location_id':loc,
            'concern_keys':u['concern_keys'],'condition':u['trigger_condition'],'mechanism':u['failure_mechanism'],
            'suggested_sections':refs,'selected':u['scenario_id'] in selected,'source_members':u['members']})
    unmapped=[{'collection':k,**u} for k in ('G','E','T') for u in bundle[k] if u['mapping_status']=='unmapped']
    lines+=['','## 未映射记录','']+[f"- {u['collection']} / {u['scenario_id']}：{u['reason']}；{refs_text(u['source_refs'])}" for u in unmapped]
    warnings=[r for r in bundle['graphs'] if r.get('warnings')]
    lines+=['','## 输入与推理诊断','']+[f"- {r['id']}：{'；'.join(r['warnings'])}" for r in warnings]
    (root/'report.md').write_text('\n'.join(lines)+'\n',encoding='utf8')
    write_json(root/'missing_report.json',{'missing':rows,'unmapped':unmapped,'scope':bundle['scope']})
    write_diagrams(root/'diagrams',bundle['ssd'],bundle['mounts'])
    # AR is visible for architectural completeness, without call arrows or mounts.
    trace=[u for group in bundle['ssd'].get('use_case_hierarchies',[]) for u in group['units'] if u['layer']=='AR']
    trace+=bundle.get('ar_trace',[])
    if trace:
        path=root/'diagrams/architecture.mmd'
        text=path.read_text(encoding='utf8')+'  subgraph AR_TRACE["AR 实现追溯（不参与异常分析）"]\n'
        for i,u in enumerate(trace):
            text+='    ar'+str(i)+'["'+str(u.get('name',u.get('source_id','AR'))).replace('"',"'")+'"]\n'
        path.write_text(text+'  end\n',encoding='utf8')
    workbook=Workbook();workbook.remove(workbook.active)
    def sheet(name,header,data):
        ws=workbook.create_sheet(name);ws.append(header)
        for row in data:ws.append(row)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24476C')
        for i in range(1,len(header)+1):ws.column_dimensions[get_column_letter(i)].width=24 if i<4 else 45
        for row in ws:
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
        return ws
    metric_rows=[]
    for bucket in [values['overall'],*values['concerns'],*values['groups'],values['unclassified']]:
        key=bucket['category'];description=defs[key]['label']+': '+defs[key]['description'] if key in defs else ('异常总计' if key=='all' else GROUP_LABELS.get(key,key))
        for m in bucket['metrics']:
            metric_rows.append([key,description,bucket['kind'],LABELS[m['metric']],m['numerator'],m['denominator'],m['value'],
                ', '.join(m['numerator_ids']),', '.join(m['denominator_ids'])])
    ws=sheet('指标',['类别','关注点说明','分类层级','指标','分子','分母','值','分子单元ID','分母单元ID'],metric_rows)
    for row in ws.iter_rows(min_row=2):row[6].number_format='0.00%'
    def unit_rows(us):
        return [[u['scenario_id'],u['use_case_id'],u['location_id'],', '.join(u['concern_keys']),u['exception_object'],u['operation'],
            u['failure_mechanism'],u['trigger_condition'],u['mapping_status'],refs_text(u['source_refs']),
            ', '.join(m['source_scenario_id'] for m in u['members']),', '.join(sorted({m['source_basis'] for m in u['members']})),
            ', '.join(u.get('source_layers',[u['requirement_layer']])),'；'.join(u['necessary_conditions']),u.get('derivation_reason',''),
            '；'.join(u.get('pending_details',[]))] for u in us]
    header=['单元ID','用例','交互位置','关注点','异常对象','操作','原因','触发条件','映射状态','章节行号','原场景ID','来源性质','来源层级','必要条件','推导依据','待确认细节']
    sheet('生成异常',header,unit_rows(bundle['G']));sheet('已有异常',header,unit_rows(bundle['E']))
    sheet('历史异常库',header,unit_rows(bundle['history']));sheet('测试GT',header,unit_rows(bundle['T']))
    sheet('缺失异常',header+['建议补充章节'],[row+[refs_text(details['suggested_sections']) or '章节定位待确认']
        for row,details in zip(unit_rows(missing),rows)])
    sheet('最终推荐',header,unit_rows([units[k] for k in sorted(selected)]))
    sheet('未映射异常',['集合']+header,[[u['collection']]+row for u in unmapped for row in unit_rows([u])])
    sheet('排名明细',['关注点','单元ID','同类历史数','相似度求和','类内排名','入选','选取方式'],
        [[c['concern_key'],r['scenario_id'],c['history_count'],r['score'],r['rank'],r['selected'],r['selection_reason']] for c in ranking['categories'] for r in c['items']])
    sheet('相似度明细',['关注点','候选单元','历史单元','余弦相似度'],
        [[c['concern_key'],r['scenario_id'],s['history_id'],s['cosine']] for c in ranking['categories'] for r in c['items'] for s in r['similarities']])
    sheet('匹配关系',['集合关系','左单元','右单元','判定','依据'],
        [[k,p['checker_scenario_id'],p['generated_scenario_id'],p['status'],p.get('reason','')] for k,a in values['assignments'].items() for p in a['selected_pairs']])
    sheet('正常可选',['场景ID','层级','类型','用例','步骤'],[[s['scenario_id'],s['requirement_layer'],s['scenario_type'],s['use_case_id'],str(s.get('steps',[]))] for s in bundle['normal']])
    sheet('关注点挂载',['挂载ID','交互','关注点','依据'],[[m['mount_id'],m['exchange_id'],m['concern_key'],m['mount_reason']] for m in bundle['mounts']])
    sheet('对象特性',['交互','特性','值','来源性质','推理依据','原文定位'],[[f['exchange_id'],t['name'],str(t['value']),t['basis'],t['rationale'],refs_text(t['source_refs'])] for f in bundle['features'] for t in f['traits']])
    sheet('图检查',['对象ID','检查器','通过','依据','诊断'],[[r['id'],r['kind'],r['valid'],r['reason'],'；'.join(r.get('warnings',[]))] for r in bundle['graphs']])
    sheet('评估状态',['项目','值'],[['状态',status],['协议',bundle['protocol']],['范围',', '.join(bundle['scope']['use_case_ids'])],
        ['GT来源','临时客户测试spec，包含标记的框架推导'],['误报率分母','GT；结果可能超过100%'],['随机种子',ranking['seed']]])
    try:workbook.save(root/'scene_assessment.xlsx')
    finally:workbook.close()
