"""Document text extraction with source-location metadata."""

from __future__ import annotations

import re
from pathlib import Path


def _plain_text(path: Path) -> dict:
    text = path.read_text(encoding="utf-8")
    lines = [{"location": f"line {i}", "text": line} for i, line in enumerate(text.splitlines(), 1)]
    return {"path": str(path), "format": path.suffix.lower().lstrip("."), "text": text, "locations": lines}


def _docx_text(path: Path) -> dict:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX extraction requires python-docx") from exc
    document = Document(str(path))
    locations = []
    for idx, paragraph in enumerate(document.paragraphs, 1):
        if paragraph.text.strip():
            locations.append({"location": f"paragraph {idx}", "text": paragraph.text})
    for table_idx, table in enumerate(document.tables, 1):
        for row_idx, row in enumerate(table.rows, 1):
            text = " | ".join(cell.text.strip() for cell in row.cells)
            if text.strip():
                locations.append({"location": f"table {table_idx}, row {row_idx}", "text": text})
    return {"path": str(path), "format": "docx", "text": "\n".join(x["text"] for x in locations), "locations": locations}


def _pdf_text(path: Path) -> dict:
    pages = []
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        for page_idx, page in enumerate(reader.pages, 1):
            pages.append({"location": f"page {page_idx}", "text": page.extract_text() or ""})
    except ImportError:
        try:
            import pdfplumber
            with pdfplumber.open(str(path)) as pdf:
                for page_idx, page in enumerate(pdf.pages, 1):
                    pages.append({"location": f"page {page_idx}", "text": page.extract_text() or ""})
        except ImportError as exc:
            raise RuntimeError("PDF extraction requires pypdf or pdfplumber") from exc
    text = "\n".join(page["text"] for page in pages).strip()
    if not text:
        raise RuntimeError("PDF contains no extractable text; OCR or a text-based PDF is required")
    return {"path": str(path), "format": "pdf", "text": text, "locations": pages}


def _xlsx_text(path: Path) -> dict:
    from openpyxl import load_workbook
    from .assessment import TEST_LABELS
    book=load_workbook(path,read_only=False,data_only=False)
    lines=['# '+path.stem];locations=[{'location':'workbook','text':lines[0]}]
    skipped=[]
    try:
        for sheet in book.worksheets:
            headers=None;header_row=None
            for row in sheet.iter_rows(min_row=1,max_row=min(30,sheet.max_row)):
                normalized=[TEST_LABELS.get(str(c.value).strip(),str(c.value).strip()) if c.value is not None else '' for c in row]
                if {'编号','场景名称','前置条件','测试步骤','预期结果','关联需求'}<=set(normalized):
                    headers=normalized;header_row=row[0].row;break
            if headers is None:
                skipped.append(sheet.title);continue
            merged={}
            for region in sheet.merged_cells.ranges:
                for row in range(max(header_row+1,region.min_row),region.max_row+1):
                    for col in range(region.min_col,region.max_col+1):merged[(row,col)]=sheet.cell(region.min_row,region.min_col)
            def add(text,row,cell_range):
                lines.append(text);locations.append({'location':f'{sheet.title}!{cell_range}',
                    'sheet':sheet.title,'row':row,'cell_range':cell_range,'text':text})
            add('## '+sheet.title,header_row,f'A{header_row}:{sheet.cell(header_row,len(headers)).coordinate}')
            for number in range(header_row+1,sheet.max_row+1):
                values=[]
                for col,label in enumerate(headers,1):
                    cell=merged.get((number,col),sheet.cell(number,col))
                    if label and cell.data_type=='f':raise ValueError(f'test spec formula requires explicit values: {sheet.title}!{cell.coordinate}')
                    value='' if cell.value is None else str(cell.value)
                    if label=='编号' and isinstance(cell.value,int) and re.fullmatch('0+',cell.number_format):
                        value=str(cell.value).zfill(len(cell.number_format))
                    values.append(value)
                if not any(v.strip() for v in values):continue
                row=dict(zip(headers,values));identifier=row.get('编号','').strip();name=row.get('场景名称','').strip()
                if not identifier or not name:raise ValueError(f'test spec missing case ID/name at {sheet.title}!{number}')
                add('### '+identifier+'：'+name,number,f'A{number}:{sheet.cell(number,len(headers)).coordinate}')
                for col,label in enumerate(headers,1):
                    if label and label not in ('编号','场景名称'):
                        value=values[col-1].replace('\r','').replace('\n','<br>')
                        add('- '+label+'：'+value,number,sheet.cell(number,col).coordinate)
                add('- 来源定位：'+sheet.title+' 第'+str(number)+'行',number,f'A{number}:{sheet.cell(number,len(headers)).coordinate}')
    finally:book.close()
    if len(lines)==1:raise ValueError('no test worksheet found; required columns: 用例编号、用例名称、前提条件、测试步骤、预期结果、对应需求')
    return {'path':str(path),'format':'xlsx','text':'\n'.join(lines),'locations':locations,
        'warnings':['non-test worksheet skipped: '+name for name in skipped]}


def extract_document(path: str | Path) -> dict:
    """Extract a supported document and return text plus source locations."""
    source = Path(path).expanduser().resolve()
    if not source.exists() or not source.is_file():
        raise FileNotFoundError(f"document not found: {source}")
    suffix = source.suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        result = _plain_text(source)
    elif suffix == ".docx":
        result = _docx_text(source)
    elif suffix == ".pdf":
        result = _pdf_text(source)
    elif suffix in {'.xlsx','.xlsm'}:
        result=_xlsx_text(source)
    else:
        raise ValueError("unsupported document format; use .md, .markdown, .txt, .docx, .pdf, or .xlsx (convert legacy .xls first)")
    # Keep original blank lines: downstream chapter citations use these line numbers.
    return result
