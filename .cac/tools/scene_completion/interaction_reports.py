"""Layered workbooks and source/SSD diagnoses. Values remain unrounded in JSON."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .interaction_protocol import baseline
from .role_reports import GROUP_LABELS
from .semantic_backend import write_json
from .sources import fingerprint


def refs_text(refs):
    return "；".join(f"{r['document']} → {' / '.join(r['heading_path'])}：{r['line_start']}-{r['line_end']}" for r in refs)


def category_description(kind, key):
    if kind == "concern":
        value = baseline()[key]
        return value["label"] + "：" + value["description"]
    if kind == "group":
        return GROUP_LABELS.get(key, key)
    return "异常总计：只比较异常条件，完整和部分按一一匹配计覆盖" if kind == "overall" else "未分类异常：尚未映射到有效关注点"


def coverage_diagnosis(generated, checked, matches, metrics, ssd, mounts):
    g = {s["scenario_id"]: s for s in generated}
    nodes = {n["node_id"]: n for n in ssd["nodes"]}
    exchanges = {e["exchange_id"]: e for e in ssd["exchanges"]}
    result = []
    missed = set(metrics["layers"]["combined"]["overall"]["missed_checker_ids"])
    for c in checked:
        links = [p for p in metrics['assignment']['selected_pairs'] if p["checker_scenario_id"] == c["scenario_id"]]
        competing = [p for p in metrics['assignment']['eligible_but_unselected_pairs'] if p['checker_scenario_id'] == c['scenario_id']]
        relevant = [m for m in mounts if m["requirement_layer"] == c["requirement_layer"] and m["use_case_id"] == c["use_case_id"] and
            exchanges[m["exchange_id"]].get("layer_use_case_id", m["use_case_id"]) == c.get("layer_use_case_id", c["use_case_id"]) and
            m["concern_key"] in c.get("concern_keys", [])]
        reason = "已识别" if c["scenario_id"] not in missed else (
            "关注点体系表达缺口，需逐条语义诊断并向用户确认" if not c.get("concern_keys") else
            "相应关注点未挂载到本层交互，需核实架构/交互或路由" if not relevant else
            "存在语义重合但一一分配未占用候选，需检查场景粒度或候选不足" if competing else
            "已有挂载仍未匹配，需核实具体化、操作对象/机制和匹配判断")
        result.append({"checker_scenario_id": c["scenario_id"], "requirement_layer": c["requirement_layer"],
            "use_case_id": c["use_case_id"], "layer_use_case_id": c.get("layer_use_case_id", c["use_case_id"]), "source_refs": c["source_refs"],
            "exception_object": c["exception_object"], "operation": c["operation"],
            "failure_mechanism": c["failure_mechanism"], "trigger_condition": c["trigger_condition"],
            "concern_keys": c.get("concern_keys", []), "covered": c["scenario_id"] not in missed,
            "matched_generated_ids": sorted({p["generated_scenario_id"] for p in links}),
            "matched_exchange_ids": sorted({g[p["generated_scenario_id"]]["exchange_id"] for p in links}),
            "relevant_mounts": relevant, "check_components": sorted({nodes[m["subject_node_id"]]["name"] for m in relevant}),
            'eligible_unselected_generated_ids': sorted({p['generated_scenario_id'] for p in competing}),
            "gap_reason": reason, "requires_confirmation": c["scenario_id"] in missed and not c.get("concern_keys")})
    return result


def write_diagrams(root, ssd, mounts):
    """Portable Mermaid sources include actual rewritten resource-service lifelines."""
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    nodes = {n["node_id"]: n for n in ssd["nodes"]}
    aliases = {nid: "n" + str(i) for i, nid in enumerate(nodes)}
    def label(text):
        return str(text).replace('"', "'").replace("\n", " ").replace(";", "；")
    architecture = ["flowchart LR"]
    used = {e[k] for e in ssd["exchanges"] for k in ("from_node", "to_node")}
    for nid in sorted(used):
        n = nodes[nid]
        architecture.append(f'  {aliases[nid]}["{label(n["name"])} ({n.get("layer", "")})"]')
    edges = set()
    for e in ssd["exchanges"]:
        key = (e["from_node"], e["to_node"], e["requirement_layer"])
        if key not in edges:
            architecture.append(f'  {aliases[key[0]]} -->|{key[2]}| {aliases[key[1]]}')
            edges.add(key)
    for m in ssd["resource_service_mappings"]:
        physical, service = m["physical_resource_id"], m["resource_service_id"]
        architecture.append(f'  {aliases[physical]}[("{label(nodes[physical]["name"])} 物理资源")]')
        architecture.append(f'  {aliases[service]} -. 物理资源映射 .-> {aliases[physical]}')
    (root / "architecture.mmd").write_text("\n".join(architecture) + "\n", encoding="utf-8")
    dependency = ["flowchart LR"]
    for e in ssd["exchanges"]:
        endpoints = [nodes[e[k]] for k in ("from_node", "to_node")]
        if e["requirement_layer"] == "SR" and all(n.get("kind") in {"internal_service", "abstract_service"} for n in endpoints):
            dependency.append(f'  {aliases[e["from_node"]]}["{label(endpoints[0]["name"])}"] -->|{label(e["operation"])}| {aliases[e["to_node"]]}["{label(endpoints[1]["name"])}"]')
    (root / "service_dependencies.mmd").write_text("\n".join(dependency) + "\n", encoding="utf-8")
    hierarchy = ["flowchart TD"]
    for group_number, group in enumerate(ssd.get("use_case_hierarchies", [])):
        units = group["units"]
        unit_aliases = {u["unit_id"]: f"u{group_number}_{i}" for i, u in enumerate(units)}
        hierarchy.append(f'  subgraph g{group_number}["{label(group["use_case_id"])}"]')
        for unit in units:
            hierarchy.append(f'    {unit_aliases[unit["unit_id"]]}["{unit["layer"]}: {label(unit["source_id"])} {label(unit["name"])}"]')
        for unit in units:
            for parent in unit["parent_ids"]:
                hierarchy.append(f'    {unit_aliases[parent]} --> {unit_aliases[unit["unit_id"]]}')
        hierarchy.append("  end")
    (root / "use_case_decomposition.mmd").write_text("\n".join(hierarchy) + "\n", encoding="utf-8")
    write_json(root / "use_case_decomposition.json", ssd.get("use_case_hierarchies", []))
    write_json(root / 'rr_business_service_roles.json', ssd.get('rr_business_service_roles', []))
    roles = ['# RR逻辑业务服务职责与SR映射', '', '逻辑职责不新增部署组件；异常沿各层实际交互独立生成。', '']
    for role in ssd.get('rr_business_service_roles', []):
        roles += [f"- {role['use_case_id']}：{role['service_type']} / RR {role['rr_exchange_id']}",
            '  对应SR：' + ', '.join(role['mapped_sr_exchange_ids']),
            '  正向依据：' + role['evidence_quote'] + '；' + refs_text(role['source_refs'])]
    (root / 'rr_business_service_roles.md').write_text('\n'.join(roles) + '\n', encoding='utf-8')
    diagram_files = []
    for uid in dict.fromkeys(e["use_case_id"] for e in ssd["exchanges"]):
        for layer in ("RR", "SR", "fused"):
            selected = [e for e in ssd["exchanges"] if e["use_case_id"] == uid and (layer == "fused" or e["requirement_layer"] == layer)]
            out = ["sequenceDiagram", "  autonumber"]
            participants = list(dict.fromkeys(e[k] for e in selected for k in ("from_node", "to_node")))
            for nid in participants:
                out.append(f'  participant {aliases[nid]} as {label(nodes[nid]["name"])}')
            exchange_lookup = {e["exchange_id"]: e for e in selected}
            for message in ssd["messages"]:
                if message["exchange_id"] not in exchange_lookup:
                    continue
                e = exchange_lookup[message["exchange_id"]]
                if message["direction"] == "request":
                    if e.get("guard"):
                        out.append(f'  Note over {aliases[e["from_node"]]}: {label(e["guard"])}')
                    arrow = "->>"
                    text = f'[{e["requirement_layer"]}] {label(message["text"])} ({e["exchange_id"]})'
                else:
                    arrow, text = "-->>", label(message["text"])
                out.append(f'  {aliases[message["from_node"]]}{arrow}{aliases[message["to_node"]]}: {text}')
                if message["direction"] == "request":
                    keys = sorted({m["concern_key"] for m in mounts if m["exchange_id"] == e["exchange_id"]})
                    out.append(f'  Note over {aliases[e["to_node"]]}: 关注点 {", ".join(keys)}')
            from .interaction_diagrams import render_interaction_ssd, diagram_stem
            stem = diagram_stem(uid, layer)
            (root / (stem + ".mmd")).write_text("\n".join(out) + "\n", encoding="utf-8")
            render_interaction_ssd(ssd, uid, layer, root / (stem + '.svg'))
            diagram_files.append({'use_case_id': uid, 'layer': layer, 'source': stem + '.mmd', 'svg': stem + '.svg'})
    write_json(root / 'ssd_diagram_files.json', diagram_files)


def export(root, generated, checked, matches, metrics, recommendations, adoption_result, diagnosis, mounts, acceptance_status=None):
    root = Path(root)
    from .semantic_backend import read_json
    saved_run = read_json(root/'run_manifest.json') if (root/'run_manifest.json').exists() else {}
    scope = saved_run.get('evaluation_scope', {'kind': 'full'})
    scope_text = ('代表用例验证：' + '、'.join(scope['use_case_ids']) + '；不代表全项目覆盖') if scope.get('kind') == 'representative' else '全输入用例范围'
    if scope.get('exchange_local_ids'):
        scope_text += '；局部交互链：' + '、'.join(scope['exchange_local_ids']) + '；不代表整个用例覆盖，原文异常分母保留'
    verification_text = '最终独立效果审查（未重复全量比较）' if saved_run.get('comparison_verification') == 'final-review' else '全量独立第二批比较及最终效果审查'
    quality_text = '保留来源质量诊断；本次仅验证结构与方法流程，不声明样例输入完整' if saved_run.get('source_quality_gate')=='structural' else '严格来源质量核验'
    result_hash = fingerprint([generated, checked, matches, metrics, adoption_result])
    accepted = bool(acceptance_status and acceptance_status.get("accepted") and acceptance_status.get("result_hash") == result_hash)
    status = "验收通过，详见 acceptance.json" if accepted else "阶段性指标，尚未通过正式验收"
    write_json(root / "assessment_status.json", {"accepted": accepted, "result_hash": result_hash, "status": status,
        'evaluation_scope': scope, 'comparison_verification': saved_run.get('comparison_verification', 'all')})
    write_json(root / "metrics.json", metrics)
    write_json(root / "coverage_diagnosis.json", diagnosis)
    write_json(root / "adoption.json", adoption_result)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "指标"
    sheet.append(["层级", "类别", "关注点说明", "统计类型", "已有异常分母", "识别漏报分子", "异常漏报率",
        "生成异常分母", "已匹配生成分子", "异常完整率", "完整覆盖已有数", "仅部分覆盖已有数", "同类挂载已有覆盖数", "漏报场景ID", "验收范围"])
    for layer, data in metrics["layers"].items():
        rows = [data["overall"], *data["concerns"].values(), *data["groups"].values(), data["unclassified"]]
        for row in rows:
            sheet.append([layer, row["category"], category_description(row["category_kind"], row["category"]),
                row["category_kind"], row["checker_count"], row["miss_count"], row["miss_rate"],
                row["generated_count"], row["overlap_count"], row["completeness"], row["full_checker_count"],
                row["partial_only_checker_count"], len(row["same_category_checker_ids"]), ", ".join(row["missed_checker_ids"]),
                '独立报告，不设漏报阈值' if layer in {'RR', 'SR'} else '两层合计报告'])
            sheet.cell(sheet.max_row, 7).number_format = "0.00%"
            sheet.cell(sheet.max_row, 10).number_format = "0.00%"
    for title, scenes in (("生成场景", generated), ("已有异常", checked)):
        table = workbook.create_sheet(title)
        table.append(["ID", "层级", "用例", "类型", "场景", "异常对象", "操作", "失败机制", "触发条件", "关注点", "SSD交换", "来源章节", "正常或可选步骤"])
        for s in scenes:
            table.append([s["scenario_id"], s["requirement_layer"], s["use_case_id"], s["scenario_type"], s.get("name"),
                s.get("exception_object"), s.get("operation"), s.get("failure_mechanism"), s.get("trigger_condition"),
                ", ".join(s.get("concern_keys", [])), s.get("exchange_id"), refs_text(s.get("source_refs", [])), " → ".join(s.get("steps", []))])
    table = workbook.create_sheet("匹配关系")
    table.append(["已有ID", "生成ID", "判定", "共同异常条件", "缺少条件细节", "依据", '一一匹配计数'])
    selected = {(p['checker_scenario_id'], p['generated_scenario_id']) for p in metrics['assignment']['selected_pairs']}
    for p in matches:
        if p["status"] != "unmatched":
            table.append([p["checker_scenario_id"], p["generated_scenario_id"], p["status"], p.get("specific_shared_condition"), "; ".join(p.get("missing_condition_details", [])), p["reason"],
                '计入' if (p['checker_scenario_id'], p['generated_scenario_id']) in selected else '未分配，不计入指标'])
    table = workbook.create_sheet("分类差异")
    table.append(["已有ID", "生成ID", "已有关注点", "生成关注点", "异常已识别"])
    for p in metrics["classification_differences"]:
        table.append([p["checker_scenario_id"], p["generated_scenario_id"], ", ".join(p["checker_concerns"]), ", ".join(p["generated_concerns"]), "是"])
    table = workbook.create_sheet("章节诊断")
    table.append(["已有ID", "层级", "来源章节", "异常条件", "关注点", "检查组件", "SSD交换", "已识别", "缺口原因"])
    for d in diagnosis:
        table.append([d["checker_scenario_id"], d["requirement_layer"], refs_text(d["source_refs"]), d["trigger_condition"], ", ".join(d["concern_keys"]), ", ".join(d["check_components"]), ", ".join(d["matched_exchange_ids"]), d["covered"], d["gap_reason"]])
    table = workbook.create_sheet("推荐来源明细")
    table.append(["ID", "层级", "用例", "支持度", "缺失度", "推荐评分", "待确认", "建议补充", "证据章节", "SSD交换", "测试对照"])
    for r in recommendations:
        table.append([r["scenario_id"], r["requirement_layer"], r["use_case_id"], r["support_score"], r["missing_score"], r["confidence"], r["needs_confirmation"], r["recommendation"], refs_text(r["source_refs"]), r["exchange_id"], '独立测试全量对照，不按RR/SR分区'])
    table = workbook.create_sheet('推荐')
    table.append(['推荐单元ID', '关注点', '异常条件', '原始RR/SR场景ID', 'SSD位置', '来源层级', '一一采纳结果'])
    for unit in (adoption_result.get('equivalence') or {}).get('units', []):
        table.append([unit['scenario_id'], ', '.join(unit['concern_keys']), unit['trigger_condition'],
            ', '.join(unit['member_scenario_ids']), ', '.join(unit['member_exchange_ids']), ', '.join(unit['source_layers']),
            '采纳' if unit['scenario_id'] in adoption_result['adopted_ids'] else '未采纳'])
    table = workbook.create_sheet("采纳结果")
    table.append(["层级", "推荐分母", "采纳分子", "采纳率", "对照说明"])
    table.append(['全部推荐', adoption_result["recommended_count"], adoption_result["adopted_count"], adoption_result["adoption_rate"], 'RR/SR同一交互关注点异常语义去重；全量测试按交互/关注点拆分，一一匹配；含明确标记的框架推导测试单元'])
    table.cell(2, 4).number_format = "0.00%"
    for layer, detail in adoption_result['recommendation_layers'].items():
        denominator, numerator = len(detail['recommended_ids']), len(detail['adopted_ids'])
        table.append([layer + '推荐来源', denominator, numerator, numerator / denominator if denominator else None,
            '按原始成员来源汇总，不将测试条件分区；跨层重复成员数量不可相加'])
        table.cell(table.max_row, 4).number_format = '0.00%'
    table.append(['原文明示测试贡献', None, adoption_result.get('explicit_adopted_count', 0), None, '原文明示原因对应的采纳单元数'])
    table.append(['框架推导测试贡献', None, adoption_result.get('derived_adopted_count', 0), None, '按已确认口径展开的可能原因，不冒充原文已写明原因'])
    tests_file = root / 'test_conditions.json'
    if tests_file.exists():
        from .semantic_backend import read_json
        table = workbook.create_sheet('测试异常单元')
        table.append(['单元ID', '原测试ID', '关注点', '交互位置', '异常条件', '来源类型', '原文条件', '推导依据', '章节'])
        for condition in read_json(tests_file)['scenarios']:
            table.append([condition['scenario_id'], condition['test_scenario_id'], ', '.join(condition['concern_keys']),
                condition['interaction_location'], condition['trigger_condition'], condition['source_basis'],
                condition['original_condition_quote'], condition.get('derivation_reason'), refs_text(condition['source_refs'])])
    table = workbook.create_sheet("正常可选统计")
    table.append(["层级", "主成功场景", "可选场景", "指标范围说明"])
    for layer, counts in metrics["non_exception_counts"].items():
        table.append([layer, counts["main_success"], counts["alternative"], "保留导出，不参与异常匹配/漏报/完整率/推荐"])
    table = workbook.create_sheet("关注点挂载")
    table.append(["挂载ID", "层级", "用例", "关注点", "关注点说明", "SSD交换", "消息", "检查组件", "方向", "来源章节"])
    for m in mounts:
        table.append([m["mount_id"], m["requirement_layer"], m["use_case_id"], m["concern_key"], baseline()[m["concern_key"]]["label"] + "：" + m["concern_description"], m["exchange_id"], ", ".join(m["message_ids"]), m["subject_node_id"], m["payload_direction"], refs_text(m["source_refs"])])
    actors_file = root / 'participant_checks.json'
    if actors_file.exists():
        from .semantic_backend import read_json
        table = workbook.create_sheet('参与者补全')
        table.append(['用例', '层级', '参与者', '正常调用依据', '角色依据', '说明', '原文修改需确认'])
        for row in read_json(actors_file)['actor_completions']:
            table.append([row['use_case_id'], row['requirement_layer'], row['participant_name'], row['interaction_quote'],
                row['role_quote'], row['reason'], '是'])
    table = workbook.create_sheet("评估状态")
    table.append(["项目", "说明"])
    table.append(["验收状态", status])
    table.append(["结果版本", result_hash])
    table.append(['验收范围', '方法流程验证；RR、SR漏报率如实报告，不设阈值；原文覆盖和测试采纳均一一匹配'])
    table.append(['分析边界', 'RR五类逻辑业务服务职责有本层正向依据及真实SR映射；独立生成与同层匹配；可取消不自动挂异常关注点'])
    table.append(['用例范围', scope_text])
    table.append(['匹配复核方式', verification_text])
    table.append(['输入质量口径', quality_text])
    if not accepted:
        table.append(["验收要求", "完整回归、同版本续跑、离线复验及 gpt-6-luna/max 效果审查通过；漏报率不是通过门槛"])
    source_audit=root/'model_source_audit.json'
    if source_audit.exists():
        from .interaction_audit import normal_model_findings
        findings=normal_model_findings(read_json(source_audit).get('items',[]))
        table=workbook.create_sheet('输入质量诊断')
        table.append(['用例','检查项','原文/交互位置','问题','原文依据','处理口径'])
        for finding in findings:
            table.append([finding['use_case_id'],finding.get('code',finding['finding_kind']),
                finding.get('line_id',finding.get('local_exchange_id')),
                '；'.join(finding.get('issues',[])) or finding.get('reason'),
                finding.get('evidence_quote'),quality_text])
    for table in workbook:
        table.freeze_panes = "A2"
        table.auto_filter.ref = table.dimensions
        for cell in table[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245A81")
        for row in table:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in table.columns:
            table.column_dimensions[column[0].column_letter].width = min(65, max(15, max(len(str(c.value or "")) for c in column[:30]) * .8))
    sheet.column_dimensions["C"].width = 58
    try:
        workbook.save(root / "scene_assessment.xlsx")
    finally:
        workbook.close()
    lines = ["# RR/SR 异常识别评估", "", f"**{status}**", "", "本轮验证方法流程，漏报率如实报告、不设置达标阈值；原文覆盖和测试采纳均一一匹配。RR业务服务职责绑定本层SSD及真实SR映射，不复制异常、不跨层覆盖。", "", "正常/可选场景保留，均不参与下列异常指标。分类差异不等于识别漏报。", "",
        scope_text + '。', '复核方式：' + verification_text + '。', '输入质量：'+quality_text+'。', '',
        "| 层级 | 已有异常 | 漏报 | 漏报率 | 生成异常 | 异常完整率 |", "| --- | ---: | ---: | ---: | ---: | ---: |"]
    for layer, data in metrics["layers"].items():
        m = data["overall"]
        rate = lambda x: "null" if x is None else f"{x:.2%}"
        lines.append(f'| {layer} | {m["checker_count"]} | {m["miss_count"]} | {rate(m["miss_rate"])} | {m["generated_count"]} | {rate(m["completeness"])} |')
    lines += ["", f'全部推荐采纳率：{rate(adoption_result["adoption_rate"])}（{adoption_result["adopted_count"]}/{adoption_result["recommended_count"]}）。推荐按真实RR/SR映射语义去重，测试不按层级分区，按交互关注点拆分后一一匹配。',
        f'原文明示测试贡献：{adoption_result.get("explicit_adopted_count", 0)}；框架推导测试贡献：{adoption_result.get("derived_adopted_count", 0)}。推导单元不是原文已写明的测试原因。',
        "", "推荐评分为未校准的支持度/缺失度加权值，不是真实正确概率。", "", "## 未覆盖异常及待确认", ""]
    for d in diagnosis:
        if not d["covered"]:
            lines.append(f'- {d["requirement_layer"]} {refs_text(d["source_refs"])} → {d["trigger_condition"]}：{d["gap_reason"]}')
    (root / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
