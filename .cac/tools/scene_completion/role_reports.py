"""Readable three-role reports and workbooks, with chapter-level gap lists."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .semantic_backend import write_json


def ref_text(ref: dict) -> str:
    return f"{ref['document']} → {' → '.join(ref['heading_path'])}（{ref['line_start']}-{ref['line_end']}行）"


def target_text(scene: dict) -> str:
    step = scene.get("source_step_index") or scene.get("anchor_step_index")
    anchor = f" → {scene.get('use_case_id', '用例未定位')} / " + (f"步骤 {step}" if step else "用例级（步骤未定位）")
    return "\n".join(f"{doc}: " + ("；".join(ref_text(r) + anchor for r in refs) if refs else "章节关联缺失")
                     for doc, refs in scene.get("target_sections", {}).items())


def rate_text(metric: dict) -> str:
    return f"{metric['numerator']}/{metric['denominator']} = {metric['rate']:.2%}" if metric["rate"] is not None else "不可计算（分母为0）"


def _cell(value):
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False)
    # Model/source content is data, never an Excel formula.
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in rows:
        ws.append([_cell(v) for v in row])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = 25 if len(headers) > 6 else 38
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def export_reports(output, generated, checker, matches, metrics, recommendations):
    root = Path(output)
    root.mkdir(parents=True, exist_ok=True)
    for name, value in (("generated_scenarios", generated), ("existing_scenarios", checker),
                        ("scenario_matches", matches), ("metrics", metrics), ("recommendations", recommendations)):
        write_json(root / f"{name}.json", value)
    overall = metrics["overall"]
    lines = ["# 三 Agent 场景评估", "", f"- 漏报率：{rate_text(overall['miss_rate'])}",
             f"- 当前已有完整率：{rate_text(overall['existing_completeness'])}",
             f"- 匹配后端：{matches['backend']}；推荐后端：{recommendations['backend']}；rerank：{recommendations['rerank']}",
             "- 完整与部分匹配均计重合；分类分别去重，数量不可直接相加。",
             "- 推荐评分未经概率校准；候选不代表已证实的产品缺陷。", "",
             "## 各关注点指标", "", "|关注点|漏报率|当前已有完整率|", "|---|---|---|"]
    coverage = metrics.get("classification_coverage", {})
    for role, values in coverage.items():
        lines.insert(7, f"- {role} 显式异常标签覆盖：{values['classified_count']}/{values['explicit_exception_count']}；"
                     f"未分类 {len(values['unclassified_ids'])} 项。未分类异常只计入总计及未分类桶，分类指标依赖标签覆盖。")
    for key, value in metrics["by_concern"].items():
        lines.append(f"|{key}（{value['label']}）|{rate_text(value['miss_rate'])}|{rate_text(value['existing_completeness'])}|")
    lines += ["", "## 各大类指标", "", "|类别|漏报率|当前已有完整率|", "|---|---|---|"]
    for key, value in metrics["by_group"].items():
        lines.append(f"|{key}|{rate_text(value['miss_rate'])}|{rate_text(value['existing_completeness'])}|")
    lines += ["", "## 主成功、可选与未分类指标", "", "|类别|漏报率|当前已有完整率|", "|---|---|---|"]
    for key, value in metrics["special_categories"].items():
        lines.append(f"|{key}|{rate_text(value['miss_rate'])}|{rate_text(value['existing_completeness'])}|")
    for title, status in (("完整匹配", "full"), ("部分匹配与缺少行为", "partial")):
        lines += ["", f"## {title}", ""]
        for m in matches["matches"]:
            if m["status"] == status:
                lines += [f"- {m['checker_scenario_id']} ↔ {m['generated_scenario_id']}：{m['status']}。"
                          f"{m['evidence']} 缺少行为：{'；'.join(m.get('missing_behavior', [])) or '无'}"]
    lines += ["", f"## 分类差异（{len(metrics['classification_discrepancies'])} 条）", "",
              "两端标签不一致可能来自一端未分类或分类口径不同；逐条记录保存在 metrics.json。"]
    lines += ["", "## 生成器漏报", ""]
    for s in overall["missing_scenarios"]:
        lines += [f"### {s['scenario_id']}：{s['name']}", "", s["trigger"], "",
                  *[ref_text(ref) for ref in s.get("source_refs", [])], ""]
    lines += ["", "## 按章节查看未匹配场景与补充建议", ""]
    for s in recommendations["items"]:
        lines += [f"### {s['scenario_id']}：{s['name']}", "",
                  f"推荐评分：{s['confidence']:.4f}；支持度：{s['support_score']:.4f}；缺失度：{s['missing_score']:.4f}；等级：{s['priority']}",
                  "", target_text(s), "", f"建议补充：{s['trigger']}", "", s["basis"], ""]
    (root / "report.md").write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    wb = Workbook()
    wb.remove(wb.active)
    metric_rows = []
    for level, rows in (("总计", {"overall": overall}), ("关注点", metrics["by_concern"]),
                        ("大类", metrics["by_group"]), ("特殊类型", metrics["special_categories"])):
        for key, row in rows.items():
            metric_rows.append([level, key, row["checker_count"], row["generated_count"],
                                row["miss_rate"]["numerator"], row["miss_rate"]["denominator"], row["miss_rate"]["rate"],
                                row["existing_completeness"]["numerator"], row["existing_completeness"]["denominator"],
                                row["existing_completeness"]["rate"]])
    _sheet(wb, "指标", ["粒度", "类别", "已有场景数", "生成场景数", "漏报分子", "漏报分母", "漏报率",
                         "完整率分子", "完整率分母", "当前已有完整率"], metric_rows)
    for row in wb["指标"].iter_rows(min_row=2):
        row[6].number_format = "0.00%"
        row[9].number_format = "0.00%"
    for title, scenes in (("生成场景", generated["scenarios"]), ("已有场景", checker["scenarios"]),
                           ("生成器漏报", overall["missing_scenarios"])):
        _sheet(wb, title, ["ID", "用例", "类型", "名称", "前置条件", "触发", "步骤", "预期结果", "恢复", "后置条件", "关注点", "证据章节", "补充章节"],
               [[s["scenario_id"], s["use_case_id"], s["scenario_type"], s.get("name"), s.get("preconditions"), s["trigger"],
                 s["scenario_steps"], s["expected_result"], s.get("recovery"), s.get("postconditions"), s.get("concern_keys"),
                 "\n".join(ref_text(r) for r in s.get("source_refs", [])), target_text(s)] for s in scenes])
    _sheet(wb, "匹配关系", ["已有场景", "生成场景", "状态", "依据", "缺少行为"],
           [[m["checker_scenario_id"], m["generated_scenario_id"], m["status"], m["evidence"], m.get("missing_behavior", [])]
            for m in matches["matches"]])
    for title, status in (("完整匹配", "full"), ("部分匹配", "partial")):
        _sheet(wb, title, ["已有场景", "生成场景", "状态", "依据", "缺少行为"],
               [[m["checker_scenario_id"], m["generated_scenario_id"], m["status"], m["evidence"], m.get("missing_behavior", [])]
                for m in matches["matches"] if m["status"] == status])
    _sheet(wb, "推荐", ["ID", "用例", "名称", "评分", "支持度", "缺失度", "等级", "依据", "补充章节", "证据", "rerank"],
           [[s["scenario_id"], s["use_case_id"], s["name"], s["confidence"], s["support_score"],
             s["missing_score"], s["priority"], s["basis"], target_text(s), s.get("evidence_refs"),
             recommendations["rerank"]] for s in recommendations["items"]])
    wb.save(root / "scene_assessment.xlsx")
    wb.close()
    return {"report": str(root / "report.md"), "workbook": str(root / "scene_assessment.xlsx")}
