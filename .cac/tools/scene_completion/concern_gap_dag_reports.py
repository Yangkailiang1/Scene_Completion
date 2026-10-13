"""Partial results never carry a whole-project denominator or recommendation ranking."""
from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment
from .semantic_backend import write_json
from .interaction_protocol import baseline
from .concern_gap_reports import GROUP_LABELS


def partial_case(root, bundle):
    root = Path(root)
    missing = set(bundle['missing_ids'])
    units = [u for u in bundle['G'] if u['scenario_id'] in missing]
    write_json(root / 'missing_report.json', {'complete_scope': 'use_case', 'project_complete': False,
        'formal_metrics': None, 'missing': units, 'relations': bundle['relations']})
    lines = ['# 用例局部缺失报告', '', '本用例集合已完成；项目历史库及正式推荐、指标尚未发布。', '',
        '| 异常ID | 交互位置 | 关注点 | 异常条件 |', '| --- | --- | --- | --- |']
    for u in units:
        lines.append('| ' + ' | '.join(str(v).replace('|', '\\|').replace('\n', ' ') for v in
            (u['scenario_id'], u['location_id'], ', '.join(u['concern_keys']), u['trigger_condition'])) + ' |')
    (root / 'report.md').write_text('\n'.join(lines) + '\n', encoding='utf8')
    wb = Workbook()
    ws = wb.active
    ws.title = '状态'
    ws.append(['范围', '用例局部结果'])
    ws.append(['项目完整', False])
    ws.append(['正式指标', '尚未发布；未完成用例不计入漏报'])
    definitions = baseline()
    for key, title in [('G', '生成异常'), ('E', '已有异常'), ('T', '临时测试GT'), ('missing', '缺失异常')]:
        sheet = wb.create_sheet(title)
        sheet.append(['ID', '类别', '关注点说明', '交互位置', '异常对象', '异常条件', '原文来源'])
        for u in units if key == 'missing' else bundle[key]:
            labels = [str(definitions.get(k, {}).get('label', k)) + '：' + str(definitions.get(k, {}).get('description', '')) for k in u['concern_keys']]
            sheet.append([u['scenario_id'], ', '.join(u['concern_keys']), ', '.join(labels), u['location_id'],
                u['exception_object'], u['trigger_condition'], str(u['source_refs'])])
        for column in 'ABCDEFG':
            sheet.column_dimensions[column].width = 34
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
    matches = wb.create_sheet('匹配关系')
    matches.append(['集合', '已有/测试ID', '生成/已有ID', '匹配状态', '依据', '条件缺口'])
    for collection, decisions in bundle['relations'].items():
        for decision in decisions:
            matches.append([collection, decision['checker_scenario_id'], decision['generated_scenario_id'],
                decision['status'], decision.get('reason', ''), str(decision.get('missing_condition_details', []))])
    normal = wb.create_sheet('主成功与可选')
    normal.append(['ID', '层级', '类型', '名称', '步骤', '来源'])
    for scenario in bundle['normal']:
        normal.append([scenario['scenario_id'], scenario['requirement_layer'], scenario['scenario_type'],
            scenario.get('name', ''), '\n'.join(scenario.get('steps', [])), str(scenario['source_refs'])])
    wb.save(root / 'scene_assessment.xlsx')
    wb.close()


def progress(root, state, selected):
    root = Path(root)
    cases = {}
    global_tasks = []
    for record in state['tasks'].values():
        if record['case'] == '@project':
            global_tasks.append({'stage': record['stage'], 'status': record['status'], 'reason': record.get('reason', record.get('failure'))})
            continue
        cases.setdefault(record['case'], []).append(record)
    rows = []
    for case, tasks in sorted(cases.items()):
        rows.append({'use_case_id': case, 'assessment_selected': case in selected,
            'complete': all(t['status'] == 'passed' for t in tasks),
            'states': dict(Counter(t['status'] for t in tasks)),
            'pending': [{'stage': t['stage'], 'status': t['status'], 'reason': t.get('reason', t.get('failure'))}
                for t in tasks if t['status'] != 'passed']})
    published = bool(global_tasks) and all(t['status'] == 'passed' for t in global_tasks)
    record = {'complete': all(r['complete'] for r in rows) and published,
        'formal_metrics': {'path': 'metrics.json'} if published else None,
        'ranking_published': published, 'cases': rows, 'global_tasks': global_tasks}
    write_json(root / 'progress.json', record)
    lines = ['# 用例执行状态', '', '未完成用例不计入漏报；完整历史库就绪前不发布推荐排序及正式指标。', '',
        '| 用例 | 是否完成 | 任务状态 |', '| --- | --- | --- |']
    for row in rows:
        lines.append(f"| {row['use_case_id']} | {row['complete']} | {row['states']} |")
    (root / 'progress.md').write_text('\n'.join(lines) + '\n', encoding='utf8')
