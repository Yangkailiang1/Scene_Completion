"""Portable artifact checks must detect stale reports and altered Excel numbers."""
from pathlib import Path
import shutil

import pytest
from openpyxl import load_workbook

from scene_completion.comparison import compare_scenes
from scene_completion.role_reports import export_reports
from tests.test_three_roles import dataset, matched


def test_snapshot_rejects_changed_report_numeric_values_and_column_layout(tmp_path, monkeypatch):
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[1] / "examples"))
    from verify_online_shopping import validate_report_artifacts
    generated, checker = dataset([]), dataset([])
    matches = matched(generated, checker, [])
    metrics = compare_scenes(generated, checker, matches)
    recommendations = {"backend": "agent", "rerank": False, "items": []}
    expected, archive = tmp_path / "expected", tmp_path / "archive"
    export_reports(expected, generated, checker, matches, metrics, recommendations)
    shutil.copytree(expected, archive)
    validate_report_artifacts(archive, expected)
    report = archive / "report.md"
    report.write_text(report.read_text(encoding="utf-8") + "\n陈旧指标", encoding="utf-8")
    with pytest.raises(ValueError, match="Markdown report"):
        validate_report_artifacts(archive, expected)
    shutil.copy2(expected / "report.md", report)
    workbook = archive / "scene_assessment.xlsx"
    wb = load_workbook(workbook)
    wb["指标"]["F2"] = 999
    wb.save(workbook)
    wb.close()
    with pytest.raises(ValueError, match="workbook cell"):
        validate_report_artifacts(archive, expected)
    shutil.copy2(expected / "scene_assessment.xlsx", workbook)
    wb = load_workbook(workbook)
    wb["指标"].column_dimensions["C"].width = 3
    wb.save(workbook)
    wb.close()
    with pytest.raises(ValueError, match="readable widths"):
        validate_report_artifacts(archive, expected)


def test_snapshot_publishes_fresh_tree_and_preserves_previous_on_failed_verification(tmp_path, monkeypatch):
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[1] / "examples"))
    import snapshot_online_shopping as publisher
    output = tmp_path / "demo"
    output.mkdir()
    (output / "evaluation.json").write_text('{}')
    (output / "old-stale-file.json").write_text('historical residue')
    def fail(prepared, run, baseline, staged):
        staged.mkdir()
        (staged / "report.md").write_text('unverified')
        raise ValueError('incomplete verification')
    monkeypatch.setattr(publisher, "build_snapshot", fail)
    with pytest.raises(ValueError, match="incomplete verification"):
        publisher.snapshot('prepared', 'run', 'baseline', output)
    assert (output / "old-stale-file.json").read_text() == 'historical residue'
    def succeed(prepared, run, baseline, staged):
        staged.mkdir()
        (staged / "evaluation.json").write_text('{}')
        (staged / "report.md").write_text('verified')
        (staged.parent / "online_shopping_model.json").write_text('{"complete":true}')
    monkeypatch.setattr(publisher, "build_snapshot", succeed)
    publisher.snapshot('prepared', 'run', 'baseline', output)
    assert (output / "report.md").read_text() == 'verified'
    assert not (output / "old-stale-file.json").exists()
    assert (tmp_path / "online_shopping_model.json").read_text() == '{"complete":true}'
    assert list(tmp_path.glob('online-shopping-snapshot-*')) == []
