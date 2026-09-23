import copy
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from tools.scene_completion.assembly import assemble_results
from tools.scene_completion.concerns import _candidate_keys, _indexed_ar_mapping, _routing_context, list_concerns, plan_concern_matrix, validate_concern_matrix
from tools.scene_completion.diagrams import render_diagrams, validate_diagram_spec
from tools.scene_completion.document_extract import extract_document
from tools.scene_completion.exporters import export_workbooks
from tools.scene_completion.knowledge import load_concern
from tools.scene_completion.schemas import ValidationFailure, validate_scene_model
from tools.scene_completion.ssd import generate_ssd_bundle, validate_ssd, write_ssd_bundle
from tools.scene_completion.svg_renderer import render_system_composition_svg


def sample_model():
    nodes = [
        {"node_id": "user", "name": "用户", "kind": "human_actor"},
        {"node_id": "display", "name": "订单展示 Service", "kind": "internal_service", "service_type": "display", "classification_status": "inferred"},
        {"node_id": "compute", "name": "订单计算 Service", "kind": "internal_service", "service_type": "compute", "classification_status": "confirmed"},
        {"node_id": "unknown", "name": "待确认 Service", "kind": "internal_service", "service_type": "unknown", "classification_status": "needs_confirmation"},
        {"node_id": "ext-service", "name": "支付服务", "kind": "external_service"},
        {"node_id": "ext-db", "name": "外部财务数据库", "kind": "external_database"},
        {"node_id": "llm", "name": "外部 LLM", "kind": "external_llm"},
        {"node_id": "int-db", "name": "订单数据库", "kind": "internal_database"},
        {"node_id": "system", "name": "订单系统", "kind": "internal_service", "layer": "RR", "service_type": "unknown"},
        {"node_id": "sr-order", "name": "OrderService", "kind": "abstract_service", "layer": "SR", "use_case_id": "UC-001"},
        {"node_id": "impl-order", "name": "createOrder", "kind": "implementation_api", "layer": "AR"},
    ]
    interactions = [
        {"interaction_id": "INT-HUMAN", "use_case_id": "UC-001", "from_node": "user", "to_node": "display", "direction": "incoming", "message": "查询订单", "api": "GET /orders", "sequence": 1, "source_step_index": 1, "source_location": "line 1"},
        {"interaction_id": "INT-EXT-SERVICE", "use_case_id": "UC-001", "from_node": "compute", "to_node": "ext-service", "direction": "outgoing", "message": "发起支付", "api": "POST /payments", "sequence": 2, "source_step_index": 2},
        {"interaction_id": "INT-EXT-CALLER", "use_case_id": "UC-001", "from_node": "ext-service", "to_node": "compute", "direction": "incoming", "message": "回调支付结果", "api": "POST /payment-callback", "sequence": 3, "source_step_index": 3},
        {"interaction_id": "INT-EXT-DB", "use_case_id": "UC-001", "from_node": "compute", "to_node": "ext-db", "direction": "outgoing", "message": "查询预算", "sequence": 4, "source_step_index": 4},
        {"interaction_id": "INT-LLM", "use_case_id": "UC-001", "from_node": "compute", "to_node": "llm", "direction": "outgoing", "message": "请求模型分析", "sequence": 5, "source_step_index": 5},
        {"interaction_id": "INT-DB", "use_case_id": "UC-001", "from_node": "compute", "to_node": "int-db", "direction": "internal", "message": "保存订单", "sequence": 6, "source_step_index": 6},
        {"interaction_id": "INT-SERVICE", "use_case_id": "UC-001", "from_node": "display", "to_node": "compute", "direction": "internal", "message": "提交计算", "sequence": 7, "source_step_index": 7},
    ]
    return {
        "version": "4",
        "project": "fixture",
        "system_name": "订单系统",
        "source": {"path": "fixture.md"},
        "system_composition": {"nodes": nodes, "edges": [{"edge_id": "EDGE-1", "from_node": "user", "to_node": "display", "relation": "calls"}]},
        "use_cases": [{
            "use_case_id": "UC-001",
            "use_case_name": "查询订单",
            "actors": ["用户"],
            "preconditions": "用户已登录。",
            "trigger": "用户进入订单页面。",
            "postconditions": "系统展示订单。",
            "main_flow": [{"step_index": i, "text": f"步骤 {i}"} for i in range(1, 8)],
            "architecture": {
                "rr": {"service_id": "rr-service-UC-001", "service_name": "Order", "abstract_api_id": "RR-ORDER"},
                "sr": {"design_use_case_id": "SRUC-UC-001-ORDER", "service_id": "sr-order", "service_name": "OrderService", "abstract_api_id": "API-ORDER"},
                "ar": [{"microservice_id": "compute", "microservice_name": "订单计算 Service", "implementation_api_id": "createOrder", "implementation_api_node_id": "impl-order", "software_interface": "POST /api/v1/orders"}]
            },
        }],
        "interactions": interactions,
    }


def diagram_spec(model):
    node_ids = [node["node_id"] for node in validate_scene_model(model)["normalized_model"]["system_composition"]["nodes"]]
    interaction_ids = [item["interaction_id"] for item in model["interactions"]]
    declarations = "\n".join(f'rectangle "{node_id}" as {node_id}' for node_id in node_ids)
    return {
        "version": "4",
        "project": "fixture",
        "system_composition_diagram": {
            "node_ids": node_ids,
            "source_location": "fixture.md",
            "puml": f"@startuml\n{declarations}\n@enduml",
        },
        "interaction_concern_diagram": {
            "interaction_ids": interaction_ids,
            "source_location": "fixture.md",
            "puml": "@startuml\nactor user\nparticipant display\nuser -> display : query\n@enduml",
        },
    }


def applicable_matrix(model):
    matrix = plan_concern_matrix(model)
    for item in matrix["items"]:
        if item["interaction_id"] == "INT-HUMAN" and item["concern_key"] == "api.data.completeness":
            item.update({"status": "applicable", "basis": "GET 请求必须包含查询条件。", "exception_types": ["必填字段缺失"]})
        if item["interaction_id"] == "INT-HUMAN" and item["concern_key"] == "common.timeout":
            item.update({"status": "applicable", "basis": "查询延迟会影响用户后续操作。", "requirement_impact": "yes", "subsequent_behavior_impact": "yes", "environment_coordination_impact": "no"})
    return matrix


class SceneCompletionV2Tests(unittest.TestCase):
    def test_prebuilt_routing_index_preserves_candidate_keys(self):
        model = validate_scene_model(sample_model(), raise_on_error=True)["normalized_model"]
        context = _routing_context(model)
        for interaction in model["interactions"]:
            self.assertEqual(_candidate_keys(model, interaction), _candidate_keys(model, interaction, context))

    def test_ar_index_preserves_original_mapping_precedence_for_two_endpoints(self):
        model = sample_model()
        model["version"] = "6"
        mappings = model["use_cases"][0]["architecture"]["ar"]
        mappings[:] = [
            {"microservice_id": "first-service", "microservice_name": "First", "implementation_api_id": "first-api", "software_interface": "GET /first", "service_type": "query_read"},
            {"microservice_id": "second-service", "microservice_name": "Second", "implementation_api_id": "second-api", "software_interface": "POST /second", "service_type": "command_write"},
        ]
        normalized = validate_scene_model(model, raise_on_error=True)["normalized_model"]
        context = _routing_context(normalized)
        selected = _indexed_ar_mapping(context, "UC-001", {"second-service", "first-service"})
        self.assertEqual(selected["microservice_id"], "first-service")

    def test_model_and_concern_routing(self):
        report = validate_scene_model(sample_model())
        self.assertTrue(report["valid"], report["errors"])
        matrix = plan_concern_matrix(sample_model())
        self.assertTrue(validate_concern_matrix(sample_model(), matrix)["valid"])
        keys = {item["concern_key"] for item in matrix["items"]}
        self.assertIn("human.authentication", keys)
        self.assertIn("api.data.completeness", keys)
        self.assertIn("external_database.query_performance", keys)
        self.assertIn("external_llm.contract", keys)
        self.assertIn("internal_database.persistence", keys)
        self.assertIn("service_relation.call_order", keys)
        self.assertIn("internal_service.compute.calculation_correctness", keys)
        self.assertIn("internal_service.display.output_completeness", keys)

    def test_concern_knowledge_is_loaded_on_demand(self):
        registry = list_concerns()
        self.assertIn("common.timeout", {item["key"] for item in registry})
        self.assertNotIn("HI_DT", {item["key"] for item in registry})
        concern = load_concern("common.timeout")
        self.assertIn("common.timeout", concern["content"])
        with self.assertRaises(ValueError):
            load_concern("HI_DT")

    def test_timeout_needs_requirement_does_not_create_exception(self):
        matrix = plan_concern_matrix(sample_model())
        result = assemble_results(sample_model(), matrix, {"findings": []})
        self.assertEqual(sum(item["scenario_type"] == "concern_derived_exception" for item in result["scenario_catalog"]), 0)
        self.assertEqual(sum(item["scenario_type"] == "main_success" for item in result["scenario_catalog"]), 1)
        self.assertTrue(any(item["concern_key"] == "common.timeout" and item["status"] == "needs_requirement" for item in result["concern_matrix"]))

    def test_assemble_applicable_findings_and_stable_ids(self):
        model = sample_model()
        matrix = applicable_matrix(model)
        findings = {"findings": [
            {"interaction_id": "INT-HUMAN", "concern_key": "api.data.completeness", "exception_desc": "缺少查询条件。", "trigger": "用户提交空查询条件。", "scenario_steps": ["用户提交请求。", "系统校验失败。"], "recovery": "用户补充条件后重试。"},
            {"interaction_id": "INT-HUMAN", "concern_key": "common.timeout", "exception_desc": "查询响应超时。", "trigger": "查询等待超过需求时限。", "scenario_steps": ["用户发起查询。", "系统等待响应。"], "recovery": "系统提示重试并返回可恢复状态。"},
        ]}
        first = assemble_results(model, matrix, findings)
        second = assemble_results(model, matrix, findings)
        self.assertEqual([item["prediction_id"] for item in first["scenario_catalog"]], [item["prediction_id"] for item in second["scenario_catalog"]])
        self.assertEqual(sum(item["scenario_type"] == "concern_derived_exception" for item in first["scenario_catalog"]), 2)
        self.assertGreater(len(first["review_items"]), 0)

    def test_invalid_matrix_and_finding_rejected(self):
        model = sample_model()
        matrix = plan_concern_matrix(model)
        matrix["items"][0]["concern_key"] = "HI_DT"
        self.assertFalse(validate_concern_matrix(model, matrix)["valid"])
        valid = applicable_matrix(model)
        with self.assertRaises(ValidationFailure):
            assemble_results(model, valid, {"findings": [{"interaction_id": "INT-HUMAN", "concern_key": "human.authentication", "exception_desc": "bad", "trigger": "bad", "scenario_steps": ["bad"], "recovery": "bad"}]})

    def test_diagram_validation_and_puml_only_render(self):
        model = sample_model()
        spec = diagram_spec(model)
        self.assertTrue(validate_diagram_spec(model, spec)["valid"])
        spec["interaction_concern_diagram"]["interaction_ids"] = ["unknown"]
        self.assertFalse(validate_diagram_spec(model, spec)["valid"])
        with tempfile.TemporaryDirectory() as tmp:
            manifest = render_diagrams(model, diagram_spec(model), Path(tmp) / "diagrams")
            self.assertEqual(manifest["status"], "rendered")
            self.assertEqual({item["kind"] for item in manifest["artifacts"]}, {"system_composition_svg", "system_composition", "interaction_concern"})
            self.assertTrue(Path(manifest["manifest"]).exists())

    def test_system_composition_rejects_arrows_and_accepts_lines(self):
        model = sample_model()
        spec = diagram_spec(model)
        declarations = spec["system_composition_diagram"]["puml"].replace("@enduml", "user -> display\n@enduml")
        spec["system_composition_diagram"]["puml"] = declarations
        report = validate_diagram_spec(model, spec)
        self.assertFalse(report["valid"])
        self.assertIn("undirected", " ".join(report["errors"]))
        spec["system_composition_diagram"]["puml"] = declarations.replace("user -> display", "user -- display")
        self.assertTrue(validate_diagram_spec(model, spec)["valid"])

    def test_system_composition_uses_boundary_anchors_and_stacked_bottom_zones(self):
        model = sample_model()
        model["system_composition"]["edges"].append({"edge_id": "EDGE-2", "from_node": "display", "to_node": "system", "relation": "connects"})
        with tempfile.TemporaryDirectory() as tmp:
            path = render_system_composition_svg(model, Path(tmp) / "system.svg")
            svg = path.read_text(encoding="utf-8")
            self.assertIn("内部资源（数据库 / 知识库）", svg)
            self.assertIn("部署硬件 / 运行环境", svg)
            self.assertLess(svg.index("内部资源（数据库 / 知识库）"), svg.index("部署硬件 / 运行环境"))
            # The user box is not connected from its center (the known center is x=161.5).
            self.assertNotIn('x1="161.5"', svg)

    def test_rr_sr_fused_ssd_and_missing_ar_mapping(self):
        model = sample_model()
        bundle = generate_ssd_bundle(model, "UC-001")
        self.assertGreater(len(bundle["rr"]["messages"]), 0)
        self.assertEqual(bundle["fused"]["layer"], "fused")
        self.assertTrue(validate_ssd(bundle["fused"], model)["valid"])
        self.assertIsInstance(bundle["fused"].get("review_items", []), list)
        self.assertTrue(any(item.get("implementation_api_id") for item in bundle["ar"]["messages"]))

    def test_concern_planning_can_use_fused_ssd_messages(self):
        model = sample_model()
        bundle = generate_ssd_bundle(model, "UC-001")
        matrix = plan_concern_matrix(model, bundle["fused"])
        self.assertTrue(matrix["items"])
        self.assertTrue(all(item.get("ssd_id") == bundle["fused"]["ssd_id"] for item in matrix["items"]))
        self.assertTrue(validate_concern_matrix(model, matrix)["valid"])

    def test_sr_ar_classification_routing_is_mapping_scoped_and_independent(self):
        model = sample_model()
        model["version"] = "6"
        uc = model["use_cases"][0]
        uc["architecture"]["sr"].update({
            "service_type": "query_retrieval", "classification_status": "confirmed",
            "classification_basis": "该用例 API 查询订单", "source_location": "spec.md:API-ORDER",
        })
        uc["architecture"]["ar"][0].update({
            "service_type": "command_write", "classification_status": "confirmed",
            "classification_basis": "该实现接口写入订单", "source_location": "api.md:createOrder",
        })
        model["interactions"].extend([
            {"interaction_id": "INT-SR-IN", "use_case_id": "UC-001", "from_node": "system", "to_node": "sr-order", "direction": "outgoing", "message": "调用订单查询", "api": "API-ORDER", "abstract_api_id": "API-ORDER", "layer": "SR", "sequence": 8, "source_step_index": 1},
            {"interaction_id": "INT-SR-OUT", "use_case_id": "UC-001", "from_node": "sr-order", "to_node": "system", "direction": "incoming", "message": "返回订单查询结果", "abstract_api_id": "API-ORDER", "layer": "SR", "sequence": 9, "source_step_index": 1},
            {"interaction_id": "INT-AR", "use_case_id": "UC-001", "from_node": "sr-order", "to_node": "compute", "direction": "internal", "message": "调用写入实现", "implementation_api_id": "createOrder", "layer": "AR", "sequence": 10, "source_step_index": 1},
        ])
        first = plan_concern_matrix(model)["items"]
        by_id = {}
        for item in first:
            by_id.setdefault(item["interaction_id"], set()).add(item["concern_key"])
        sr_expected = {key for key in by_id["INT-SR-IN"] if key.startswith("sr_service.query_retrieval.")}
        self.assertEqual(len(sr_expected), 3)
        self.assertTrue({key for key in by_id["INT-SR-OUT"] if key.startswith("sr_service.query_retrieval.")})
        self.assertFalse(any(key.startswith("ar_service.") for key in by_id["INT-SR-IN"]))
        self.assertTrue(any(key.startswith("ar_service.command_write.") for key in by_id["INT-AR"]))

        changed_ar = copy.deepcopy(model)
        changed_ar["use_cases"][0]["architecture"]["ar"][0]["service_type"] = "query_read"
        second = plan_concern_matrix(changed_ar)["items"]
        second_by_id = {}
        for item in second:
            second_by_id.setdefault(item["interaction_id"], set()).add(item["concern_key"])
        self.assertEqual(sr_expected, {key for key in second_by_id["INT-SR-IN"] if key.startswith("sr_service.query_retrieval.")})
        self.assertTrue(any(key.startswith("ar_service.query_read.") for key in second_by_id["INT-AR"]))
        self.assertFalse(any(key.startswith("ar_service.command_write.") for key in second_by_id["INT-AR"]))

    def test_unknown_service_classification_is_reviewed_but_has_no_specialized_keys(self):
        model = sample_model()
        model["version"] = "6"
        model["interactions"].append({"interaction_id": "INT-SR", "use_case_id": "UC-001", "from_node": "system", "to_node": "sr-order", "direction": "outgoing", "message": "访问 SR", "abstract_api_id": "API-ORDER", "layer": "SR", "sequence": 8, "source_step_index": 1})
        normalized = validate_scene_model(model, raise_on_error=True)["normalized_model"]
        self.assertTrue(any(item["type"] == "sr_service_classification" and item["use_case_id"] == "UC-001" for item in normalized["review_items"]))
        keys = {item["concern_key"] for item in plan_concern_matrix(model)["items"] if item["interaction_id"] == "INT-SR"}
        self.assertFalse(any(key.startswith("sr_service.") for key in keys))
        self.assertFalse(any(key.startswith("ar_service.") for key in keys))

    def test_legacy_compute_never_becomes_a_new_sr_or_ar_category(self):
        model = sample_model(); model["version"] = "6"
        model["use_cases"][0]["architecture"]["sr"].update({"service_type": "compute", "classification_status": "confirmed"})
        model["use_cases"][0]["architecture"]["ar"][0].update({"service_type": "compute", "classification_status": "confirmed"})
        normalized = validate_scene_model(model, raise_on_error=True)["normalized_model"]
        sr = normalized["use_cases"][0]["architecture"]["sr"]
        ar = normalized["use_cases"][0]["architecture"]["ar"][0]
        self.assertEqual((sr["service_type"], sr["classification_status"]), ("unknown", "needs_confirmation"))
        self.assertEqual((ar["service_type"], ar["classification_status"]), ("unknown", "needs_confirmation"))
        self.assertTrue(any(item["type"] == "sr_service_classification" for item in normalized["review_items"]))
        self.assertTrue(any(item["type"] == "ar_service_classification" for item in normalized["review_items"]))

    def test_same_sr_service_name_can_have_use_case_specific_categories(self):
        model = sample_model()
        model["version"] = "6"
        first_uc = model["use_cases"][0]
        first_uc["architecture"]["sr"].update({
            "service_type": "query_retrieval", "classification_status": "confirmed",
            "classification_basis": "本 API 查询订单", "source_location": "spec.md:query",
        })
        model["interactions"].append({"interaction_id": "INT-SR-UC1", "use_case_id": "UC-001", "from_node": "system", "to_node": "sr-order", "direction": "outgoing", "message": "查询订单", "abstract_api_id": "API-ORDER", "layer": "SR", "sequence": 8, "source_step_index": 1})
        second_uc = copy.deepcopy(first_uc)
        second_uc.update({"use_case_id": "UC-002", "use_case_name": "修改订单", "main_flow": [{"step_index": 1, "text": "修改订单"}]})
        second_uc["architecture"]["rr"].update({"service_id": "rr-service-UC-002", "abstract_api_id": "RR-ORDER-UPDATE"})
        second_uc["architecture"]["sr"].update({
            "service_id": "sr-order-2", "design_use_case_id": "SRUC-UC-002", "abstract_api_id": "API-ORDER-UPDATE",
            "service_type": "resource_mutation", "classification_status": "confirmed",
            "classification_basis": "本 API 更新订单", "source_location": "spec.md:update",
        })
        model["use_cases"].append(second_uc)
        model["system_composition"]["nodes"].append({"node_id": "sr-order-2", "name": "OrderService", "kind": "abstract_service", "layer": "SR", "use_case_id": "UC-002"})
        model["interactions"].append({"interaction_id": "INT-SR-UC2", "use_case_id": "UC-002", "from_node": "system", "to_node": "sr-order-2", "direction": "outgoing", "message": "更新订单", "abstract_api_id": "API-ORDER-UPDATE", "layer": "SR", "sequence": 8, "source_step_index": 1})
        matrix = plan_concern_matrix(model)["items"]
        first_keys = {x["concern_key"] for x in matrix if x["interaction_id"] == "INT-SR-UC1"}
        second_keys = {x["concern_key"] for x in matrix if x["interaction_id"] == "INT-SR-UC2"}
        self.assertTrue(any(key.startswith("sr_service.query_retrieval.") for key in first_keys))
        self.assertTrue(any(key.startswith("sr_service.resource_mutation.") for key in second_keys))

    def test_all_sr_and_ar_taxonomy_values_have_independent_registry_routes(self):
        sr_categories = ("display_interaction", "query_retrieval", "resource_mutation", "analysis_generation", "release_activation")
        ar_categories = ("query_read", "command_write", "orchestration", "integration_event", "publish_activation")
        for index, category in enumerate(sr_categories, 1):
            model = sample_model(); model["version"] = "6"
            model["use_cases"][0]["architecture"]["sr"].update({"service_type": category, "classification_status": "confirmed", "classification_basis": "基于该 API 的业务职责", "source_location": "spec.md:sr"})
            model["interactions"].append({"interaction_id": f"INT-SR-{index}", "use_case_id": "UC-001", "from_node": "system", "to_node": "sr-order", "message": "SR 调用", "abstract_api_id": "API-ORDER", "layer": "SR", "sequence": 8, "source_step_index": 1})
            keys = {item["concern_key"] for item in plan_concern_matrix(model)["items"] if item["interaction_id"] == f"INT-SR-{index}"}
            self.assertTrue(any(key.startswith(f"sr_service.{category}.") for key in keys), category)
            self.assertFalse(any(key.startswith("ar_service.") for key in keys), category)
        for index, category in enumerate(ar_categories, 1):
            model = sample_model(); model["version"] = "6"
            model["use_cases"][0]["architecture"]["sr"].update({"service_type": "display_interaction", "classification_status": "confirmed", "classification_basis": "基于 SR API 展示职责", "source_location": "spec.md:sr"})
            model["use_cases"][0]["architecture"]["ar"][0].update({"service_type": category, "classification_status": "confirmed", "classification_basis": "基于实现 API 技术职责", "source_location": "api.md:ar"})
            model["interactions"].append({"interaction_id": f"INT-AR-{index}", "use_case_id": "UC-001", "from_node": "system", "to_node": "compute", "message": "AR 调用", "implementation_api_id": "createOrder", "layer": "AR", "sequence": 8, "source_step_index": 1})
            keys = {item["concern_key"] for item in plan_concern_matrix(model)["items"] if item["interaction_id"] == f"INT-AR-{index}"}
            self.assertTrue(any(key.startswith(f"ar_service.{category}.") for key in keys), category)
            self.assertFalse(any(key.startswith("sr_service.") for key in keys), category)

    def test_ssd_bundle_writes_four_layers_per_use_case(self):
        model = sample_model()
        bundle = generate_ssd_bundle(model, "UC-001")
        with tempfile.TemporaryDirectory() as tmp:
            manifest = write_ssd_bundle(bundle, model, Path(tmp) / "diagrams")
            self.assertEqual(set(manifest["artifacts"]), {"rr", "sr", "ar", "fused"})
            for artifact in manifest["artifacts"].values():
                self.assertTrue(Path(artifact["puml"]).exists())
                self.assertTrue(Path(artifact["json"]).exists())

    def test_extract_and_export_workbooks(self):
        model = sample_model()
        matrix = applicable_matrix(model)
        findings = {"findings": [{"interaction_id": "INT-HUMAN", "concern_key": "api.data.completeness", "exception_desc": "缺少字段。", "trigger": "请求体为空。", "scenario_steps": ["提交请求。", "拒绝请求。"], "recovery": "补充字段。"}]}
        bundle = assemble_results(model, matrix, findings)
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "source.md"
            source.write_text("# Requirements\n\nUser opens the order page.", encoding="utf-8")
            self.assertIn("User opens", extract_document(source)["text"])
            paths = export_workbooks(bundle, tmp_path / "out")
            self.assertTrue(Path(paths["prediction_workbook"]).exists())
            self.assertTrue(Path(paths["scenario_workbook"]).exists())
            self.assertTrue(Path(paths["concern_matrix"]).exists())
            mapping = load_workbook(paths["mapping_workbook"], read_only=True)
            sr_headers = [mapping["SR接口映射"].cell(3, col).value for col in range(1, mapping["SR接口映射"].max_column + 1)]
            ar_headers = [mapping["AR软件实现接口映射"].cell(3, col).value for col in range(1, mapping["AR软件实现接口映射"].max_column + 1)]
            self.assertIn("SR功能分类", sr_headers)
            self.assertIn("分类依据", sr_headers)
            self.assertIn("AR技术职责分类", ar_headers)
            self.assertIn("分类状态", ar_headers)
            workbook = load_workbook(paths["scenario_workbook"], read_only=True)
            self.assertEqual(workbook.sheetnames, ["场景清单", "关注点矩阵", "超时判断"])

    def test_v3_complete_scenarios_and_actor_separation(self):
        model = sample_model()
        model["version"] = "4"
        model["use_cases"][0]["scenarios"] = [{
            "scenario_id": "UC-001-main",
            "scenario_type": "main",
            "anchor_step_index": 0,
            "steps": model["use_cases"][0]["main_flow"],
        }, {
            "scenario_id": "UC-001-2.a",
            "scenario_type": "requirement_exception",
            "anchor_step_index": 2,
            "anchor_label": "2.a",
            "steps": [{"step_index": 1, "text": "用户查询订单"}, {"step_index": 2, "text": "系统返回订单不存在"}],
            "trigger": "订单不存在",
            "expected_result": "返回订单不存在",
            "recovery": "返回订单列表",
        }]
        report = validate_scene_model(model)
        self.assertTrue(report["valid"], report["errors"])
        matrix = applicable_matrix(model)
        findings = {"findings": [{"interaction_id": "INT-HUMAN", "concern_key": "api.data.completeness", "exception_type": "空请求", "exception_desc": "请求体为空。", "trigger": "用户提交空请求。", "source_step_index": 1, "scenario_steps": ["用户提交请求。", "系统拒绝请求。"], "recovery": "补充字段后重试。"}]}
        bundle = assemble_results(model, matrix, findings)
        self.assertEqual(sum(item["scenario_type"] == "main_success" for item in bundle["scenario_catalog"]), 1)
        self.assertTrue(any(item["scenario_type"] == "requirement_exception" for item in bundle["scenario_catalog"]))
        self.assertTrue(all(item.get("actor") != "订单系统" for item in bundle["scenario_catalog"] if item.get("actor")))

    def test_v3_ssd_has_pairs_and_hides_footbox(self):
        model = sample_model()
        model["version"] = "4"
        model["use_cases"][0]["scenarios"] = [{"scenario_id": "UC-001-main", "scenario_type": "main", "anchor_step_index": 0, "steps": model["use_cases"][0]["main_flow"]}]
        bundle = generate_ssd_bundle(model, "UC-001")
        fused = bundle["fused"]
        self.assertTrue(validate_ssd(fused, model)["valid"])
        self.assertTrue(any(item["message_kind"] == "response" for item in fused["messages"]))
        self.assertTrue(any(item["message_kind"] == "feedback" for item in bundle["rr"]["messages"]))
        puml = __import__("tools.scene_completion.ssd", fromlist=["ssd_to_puml"]).ssd_to_puml(fused, model)
        self.assertIn("hide footbox", puml)
        self.assertNotIn("Delta在线商城系统\" as", puml.split("actor ", 1)[-1] if "actor " in puml else "")


if __name__ == "__main__":
    unittest.main()
