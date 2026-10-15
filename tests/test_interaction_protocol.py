"""Behavioral guards for the new RR/SR identification evaluation units."""
import copy
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scene_completion.interaction_protocol import (PROTOCOL, baseline, build_ssd,
    route_mounts, validate_generated, identification, normalize_resource_services)
from scene_completion.interaction_batches import generation_validator, pair_validator, pair_packets, checker_validator
from scene_completion.interaction_metrics import compare, adoption
from scene_completion.interaction_reports import export
from scene_completion.sources import index_sources, validate_refs
from scene_completion.semantic_backend import read_json


@pytest.fixture
def sources(tmp_path):
    path = tmp_path / "rr.md"
    path.write_text("# Requirements\n## UC-1\n用户提交商品编号。商品编号必须存在。\n", encoding="utf-8")
    index = index_sources([str(path)])
    index["documents"][0]["requirement_layer"] = "RR"
    return index


def reference(index):
    return validate_refs([{"document": "rr.md", "line_start": 3, "line_end": 3}], index)


def test_rr_boundary_uses_human_kind_and_still_rejects_a_different_ui(sources):
    value = model(sources)
    value['nodes'][0]['boundary_role'] = 'ui'
    ssd = build_ssd(value)
    assert next(e for e in ssd['exchanges'] if e['local_id'] == 'sr1')['from_node'] == 'ui'
    value['nodes'].append({'node_id': 'other-ui', 'name': 'Other UI', 'kind': 'frontend_ui',
        'layer': 'RR', 'boundary_role': 'ui'})
    value['use_cases'][0]['exchanges'][1]['from_node'] = 'other-ui'
    with pytest.raises(ValueError, match='SR boundary differs'):
        build_ssd(value)


def model(index):
    refs = reference(index)
    return {"protocol": PROTOCOL, "source_hash": "hash", "nodes": [
        {"node_id": "actor", "name": "顾客", "kind": "human_actor", "layer": "RR"},
        {"node_id": "ui", "name": "商城UI", "kind": "frontend_ui", "boundary_role": "ui", "layer": "RR"},
        {"node_id": "service", "name": "商品服务", "kind": "internal_service", "layer": "SR", "service_types": ["query_retrieval"]},
        {"node_id": "db", "name": "内部数据库", "kind": "internal_database", "layer": "AR"},
        {"node_id": "extdb", "name": "外部数据库", "kind": "external_database", "layer": "AR"}],
        "use_cases": [{"use_case_id": "UC-1", "layer_use_case_ids": {"RR": "UC-1", "SR": "UC-1-API"}, "exchanges": [
            {"local_id": "rr1", "requirement_layer": "RR", "from_node": "actor", "to_node": "ui", "operation": "查看商品", "source_refs": refs},
            {"local_id": "sr1", "parent_rr_local_id": "rr1", "requirement_layer": "SR", "from_node": "ui", "to_node": "service", "operation": "查询商品编号", "request_fields": ["productId"], "source_refs": refs},
            {"local_id": "sr2", "parent_rr_local_id": "rr1", "parent_sr_local_id": "sr1", "requirement_layer": "SR", "from_node": "service", "to_node": "db", "operation": "读取商品", "source_refs": refs},
            {"local_id": "sr3", "parent_rr_local_id": "rr1", "parent_sr_local_id": "sr1", "requirement_layer": "SR", "from_node": "service", "to_node": "extdb", "operation": "读取外部商品", "source_refs": refs}]}]}


def scene(sid, layer="RR", keys=None, kind="exception"):
    return {"scenario_id": sid, "requirement_layer": layer, "use_case_id": "UC-1", "scenario_type": kind,
        "name": "商品不存在", "exception_object": "商品", "operation": "查询商品", "failure_mechanism": "资源不存在",
        "trigger_condition": "商品编号没有对应资源", "necessary_conditions": [], "concern_keys": keys or []}


def pair(cid, gid, status="full"):
    return {"checker_scenario_id": cid, "generated_scenario_id": gid, "status": status,
        "same_object": True, "same_operation": True, "same_mechanism": True, "compatible_conditions": True,
        "specific_shared_condition": "商品编号没有对应资源", "generated_evidence": "商品编号没有对应资源",
        "checker_evidence": "商品编号没有对应资源", "missing_condition_details": [], "reason": "对象、查询与不存在条件一致",
        "same_causal_facet": True, "causal_facet_evidence": {"generated": "资源不存在", "checker": "资源不存在", "reason": "相同具体不存在原因"},
        "identification_evidence": {k: {"generated": scene("G")[k], "checker": scene("C")[k]} for k in ("exception_object", "operation", "failure_mechanism", "trigger_condition")}}


def test_all_mandatory_concern_groups_present():
    keys = baseline()
    assert "service.display_interaction.device_adaptation" in keys
    assert sum(k.startswith("api.data.") for k in keys) == 7
    assert sum(k.startswith("internal_database.") for k in keys) == 10
    assert sum(k.startswith("service_relation.") for k in keys) == 6
    assert sum(k.startswith("external_llm.quality.") for k in keys) == 5
    assert all(not keys[k]["draft"] for k in keys)


def test_rr_one_to_multiple_sr_actual_resource_lifelines(sources):
    ssd = build_ssd(model(sources))
    assert len(ssd["resource_service_mappings"]) == 2
    assert {m["origin"] for m in ssd["resource_service_mappings"]} == {"internal", "external"}
    assert not {"db", "extdb"} & {m[k] for m in ssd["messages"] for k in ("from_node", "to_node")}
    sr = [e for e in ssd["exchanges"] if e["requirement_layer"] == "SR"]
    assert len({e["parent_exchange_id"] for e in sr}) == 1
    mounts = route_mounts(ssd)
    assert any(m["concern_key"] == "internal_database.uniqueness" for m in mounts)
    assert any(m["concern_key"] == "external_service.contract" and m["exchange_id"] == sr[-1]["exchange_id"] for m in mounts)
    assert any(m["concern_key"] == "service_relation.call_order" for m in mounts)
    assert all(m["requirement_layer"] in {"RR", "SR"} for m in mounts)


def test_db_wrapper_reuses_existing_service():
    nodes = [{"node_id": "db", "name": "DB", "kind": "internal_database"},
        {"node_id": "resource", "name": "DB资源服务", "kind": "internal_service", "physical_resource_id": "db", "resource_origin": "internal"}]
    output, mappings = normalize_resource_services(nodes)
    assert len(output) == 2 and mappings[0]["resource_service_id"] == "resource"


def test_ar_internal_calls_cannot_be_relabelled_sr(sources):
    value = model(sources)
    value["nodes"][2]["layer"] = "AR"
    with pytest.raises(ValueError, match="AR implementation"):
        build_ssd(value)


def test_external_callback_requires_receiving_boundary(sources):
    value = model(sources)
    value["nodes"][0]["kind"] = "external_service"
    value["nodes"][1]["boundary_role"] = "receiver"
    ssd = build_ssd(value)
    keys = {m["concern_key"] for m in route_mounts(ssd) if m["requirement_layer"] == "RR"}
    assert {"external_service.permission", "external_service.identity"} <= keys
    assert "human.authentication" not in keys


def test_direct_constraint_bypass_rejected(sources):
    value = model(sources)
    value["generation_constraints"] = [{"trigger": "商品不存在"}]
    with pytest.raises(ValueError, match="forbidden"):
        build_ssd(value)


def test_generated_requires_every_real_mount_and_correct_layer(sources):
    ssd = build_ssd(model(sources))
    mounts = route_mounts(ssd)
    with pytest.raises(ValueError, match="every valid mount"):
        validate_generated([], ssd, mounts, sources)
    m = mounts[0]
    value = {**scene("G"), "generation_source": "ssd_concern", "mount_id": m["mount_id"],
        "exchange_id": m["exchange_id"], "message_ids": m["message_ids"], "concern_keys": [m["concern_key"]], "source_refs": reference(sources)}
    value["requirement_layer"] = "AR"
    with pytest.raises(ValueError, match="same-layer"):
        validate_generated([value], ssd, mounts, sources)


def test_response_and_recovery_not_visible_to_matcher(tmp_path):
    g, c = scene("G"), scene("C")
    g.update({"expected_result": "SUCCESS", "recovery": "none"})
    c.update({"expected_result": "ERROR", "recovery": "retry"})
    manifest = pair_packets([g], [c], tmp_path)
    packet = read_json(tmp_path / "packets" / (manifest["batches"][0]["batch_id"] + ".json"))
    assert "expected_result" not in packet["input"]["generated"][0]
    assert "recovery" not in packet["input"]["checker"][0]
    assert pair_validator(packet, {"pairs": [pair("C", "G")]})[0]["status"] == "full"


@pytest.mark.parametrize("bad", ["same_object", "same_operation", "same_mechanism", "compatible_conditions"])
def test_incompatible_identification_cannot_match(bad):
    packet = {"input": {"generated": [scene("G")], "checker": [scene("C")]}}
    p = pair("C", "G")
    p[bad] = False
    with pytest.raises(ValueError, match="incompatible"):
        pair_validator(packet, {"pairs": [p]})


def test_partial_requires_condition_detail_not_just_topic():
    packet = {"input": {"generated": [scene("G")], "checker": [scene("C")]}}
    with pytest.raises(ValueError, match="partial requires"):
        pair_validator(packet, {"pairs": [pair("C", "G", "partial")]})


def test_one_to_one_assigns_unique_scenes_and_class_difference_is_not_miss():
    c = [scene("C1", keys=["api.data.format", "api.data.legality"]), scene("C2")]
    g = [scene("G1", keys=["service.query_retrieval.resource_existence"]), scene("G2"), scene("N", kind="main_success"), scene("A", kind="alternative")]
    pairs = [pair(cid, gid) for cid in ("C1", "C2") for gid in ("G1", "G2")]
    metrics = compare(g, c, pairs)
    rr = metrics["layers"]["RR"]
    assert rr["overall"]["checker_count"] == 2 and rr["overall"]["generated_count"] == 2
    assert rr["overall"]["miss_count"] == 0 and rr["overall"]["overlap_count"] == 2
    assert rr["concerns"]["api.data.format"]["miss_count"] == 0
    assert rr["concerns"]["api.data.format"]["same_category_checker_ids"] == []
    assert len(metrics["classification_differences"]) == 1
    assert len(metrics["assignment"]["selected_pairs"]) == 2
    assert len(metrics["assignment"]["eligible_but_unselected_pairs"]) == 2
    assert metrics["non_exception_counts"]["RR"] == {"main_success": 1, "alternative": 1}
    assert metrics["coverage_threshold_met"] is None
    assert metrics['layers']['SR']['overall']['miss_rate'] is None


def test_incomplete_or_cross_layer_pair_universe_rejected():
    with pytest.raises(ValueError, match="incomplete"):
        compare([scene("G")], [scene("C")], [])
    with pytest.raises(ValueError, match="unrelated"):
        compare([scene("G", "SR")], [scene("C", "RR")], [pair("C", "G")])
    with pytest.raises(ValueError, match="cannot publish"):
        compare([], [], [], complete=False)


def test_five_percent_is_reported_without_a_pass_threshold():
    g = [scene(f"G-{layer}-{i}", layer) for layer in ("RR", "SR") for i in range(19)]
    c = [scene(f"C-{layer}-{i}", layer) for layer in ("RR", "SR") for i in range(20)]
    pairs = [pair(cs["scenario_id"], gs["scenario_id"], "full" if cs["scenario_id"].split("-", 1)[1] == gs["scenario_id"].split("-", 1)[1] else "unmatched") for cs in c for gs in g if cs["requirement_layer"] == gs["requirement_layer"]]
    metrics = compare(g, c, pairs)
    assert metrics["coverage_threshold_met"] is None
    assert metrics['threshold'] is None
    assert metrics["layers"]["RR"]["overall"]["miss_rate"] == .05


def test_high_rr_misses_remain_visible_without_an_acceptance_gate():
    g = [scene('G-SR', 'SR')]
    c = [scene('C-SR', 'SR'), *[scene('C-RR-' + str(i)) for i in range(30)]]
    metrics = compare(g, c, [pair('C-SR', 'G-SR')])
    assert metrics['coverage_threshold_met'] is None
    assert metrics['evaluation_policy']['acceptance_layers'] == ['RR', 'SR']
    assert metrics['evaluation_policy']['cross_layer_matching'] is False
    assert metrics['layers']['RR']['overall']['miss_count'] == 30
    assert metrics['layers']['combined']['overall']['miss_rate'] == 30 / 31
    assert len(metrics['layers']['combined']['overall']['checker_ids']) == 31
    assert compare(g, [scene('C-SR', 'SR')], [pair('C-SR', 'G-SR', 'unmatched')])['coverage_threshold_met'] is None
    assert compare([], [], [])['coverage_threshold_met'] is None


def test_adoption_all_recommendations_unique_ids():
    recommended = [scene("R1", "SR"), scene("R2", "SR")]
    tests = [scene("T1", "SR"), scene("T2", "SR")]
    pairs = [pair(t["scenario_id"], r["scenario_id"], "full" if r["scenario_id"] == "R1" else "unmatched") for t in tests for r in recommended]
    result = adoption(recommended, tests, pairs)
    assert result["adoption_rate"] == .5 and result["adopted_count"] == 1
    assert adoption([], tests, [])["adoption_rate"] is None
    assert adoption([scene('RR')], [], [])['recommended_count'] == 1
    with pytest.raises(ValueError, match="recommended RR/SR exceptions"):
        adoption([scene('normal', kind='main_success')], [], [])


def test_test_adoption_has_no_layer_or_use_case_candidate_filter(tmp_path):
    generated = [scene('R-RR'), scene('R-SR', 'SR')]
    tests = [scene('T-A', 'TEST'), scene('T-B', 'TEST')]
    tests[1]['use_case_id'] = 'Different-source-number'
    manifest = pair_packets(generated, tests, tmp_path, 'interaction-adoption')
    packets = [read_json(tmp_path / 'packets' / (b['batch_id'] + '.json')) for b in manifest['batches']]
    assert len(packets) == 2 and all(p['input']['comparison_scope'] == 'test_all' for p in packets)
    assert all({s['scenario_id'] for s in p['input']['generated']} == {'R-RR', 'R-SR'} for p in packets)
    pairs = [pair(t['scenario_id'], g['scenario_id'], 'unmatched' if t['scenario_id'] == 'T-B' else 'full') for t in tests for g in generated]
    result = adoption(generated, tests, pairs)
    assert result['recommended_count'] == 2 and result['adopted_count'] == 1
    assert result['adoption_rate'] == .5 and not result['test_layer_partitioned']
    with pytest.raises(ValueError, match='incomplete'):
        adoption(generated, tests, pairs[:-1])
    # This exhaustive test scope must never leak into system/design coverage.
    packet = {**packets[0], 'stage': 'interaction-matching'}
    with pytest.raises(ValueError, match='reserved for test adoption'):
        pair_validator(packet, {'pairs': []})


def test_workbook_description_and_percent_columns_not_shifted(tmp_path):
    g, c = [scene("G")], [scene("C")]
    metrics = compare(g, c, [pair("C", "G")])
    export(tmp_path, g, c, [pair("C", "G")], metrics, [], adoption([], [], []), [], [])
    book = load_workbook(tmp_path / "scene_assessment.xlsx", read_only=True)
    try:
        sheet = book["指标"]
        assert sheet.cell(1, 2).value == "类别" and sheet.cell(1, 3).value == "关注点说明"
        assert sheet.cell(2, 5).value == 1 and sheet.cell(2, 6).value == 0
        assert sheet.cell(2, 7).value == 0 and sheet.cell(2, 10).value == 1
        assert sheet.cell(2, 7).number_format == "0.00%"
        assert book["采纳结果"].cell(2, 1).value == '全部推荐'
        assert '不将测试条件分区' in book['采纳结果'].cell(3, 5).value
        assert book["评估状态"].cell(2, 2).value == "阶段性指标，尚未通过正式验收"
        assert sheet.cell(1, 15).value == '验收范围'
        assert sheet.cell(2, 15).value == '独立报告，不设漏报阈值'
        assert '不设阈值' in book['评估状态'].cell(4, 2).value
        assert '不设置达标阈值' in (tmp_path / 'report.md').read_text(encoding='utf-8')
    finally:
        book.close()
