"""Exception-only acceptance, generation contributions, and readable exports."""
import copy
import json

import pytest
from openpyxl import load_workbook

from scene_completion.comparison import compare_scenes
from scene_completion.role_reports import export_reports
from scene_completion.sources import fingerprint
from tests.test_three_roles import dataset, link, matched, scene


def test_recommendation_targets_follow_requirement_only_assessment_scope():
    from scene_completion.three_roles import finish_recommendations
    from scene_completion.role_reports import target_text
    req = {"document": "req.md", "heading_path": ["Requirements", "UC-001"], "line_start": 4, "line_end": 5}
    design = {**req, "document": "design.md", "heading_path": ["Design", "UC-001"]}
    candidate = {**scene("G1"), "target_sections": {"req.md": [req], "design.md": [design]}}
    report = {"input_hash": "fixed", "overall": {"unmatched_generated_scenarios": [candidate]}}
    scores = [{"scenario_id": "G1", "support_score": .5, "missing_score": 1.}]
    index = {"documents": [{"document": "req.md", "role": "requirement"}, {"document": "design.md", "role": "design"}]}
    item = finish_recommendations(report, scores, "agent", False, True, index)["items"][0]
    assert item["recommendation_target_sections"] == {"req.md": [req]}
    assert item["related_design_sections"] == {"design.md": [design]}
    text = target_text(item)
    assert "建议补充（检查范围）：req.md" in text and "设计证据关联：design.md" in text
    assert "建议补充（检查范围）：design.md" not in text
    assert item["confidence"] == .7 * .5 + .3


def test_recovery_cannot_reinstate_native_rejected_pairs(tmp_path):
    from scene_completion.matching_audit import exclude_native_rejections
    from scene_completion.semantic_backend import write_json
    write_json(tmp_path / "manifest.json", {"batches": [{"batch_id": "b"}]})
    write_json(tmp_path / "results" / "b.json", {"semantic_verification": {
        "provider": "native_subagent", "proposed_matches": [link("C1", "G1"), link("C1", "G2")],
        "decisions": [{"pair_index": 0, "status": "unmatched"}, {"pair_index": 1, "status": "partial"}]}})
    candidates = [link("C1", "G1"), link("C1", "G2"), link("C1", "G3")]
    assert exclude_native_rejections(candidates, tmp_path) == candidates[1:]


def fixture_scenes():
    keys = ["common.timeout", "external_service.availability"]
    generated = dataset([
        {**scene("E1", keys), "generation_status": "explicit"},
        {**scene("E2", keys), "generation_status": "explicit"},
        {**scene("D1", keys, "concern_derived_exception"), "generation_status": "unreviewed_candidate",
         "candidate_id": "CAND-1", "subject_node_id": "NODE-1", "exchange_id": "EX-1"},
        {**scene("D2", keys, "concern_derived_exception"), "generation_status": "unreviewed_candidate"},
        scene("GM", keys, "main_success"), scene("GA", kind="alternative"),
    ])
    checker = dataset([scene("C1", keys), scene("C2", keys), scene("C3", keys),
                       scene("CM", keys, "main_success"), scene("CA", kind="alternative")])
    for item in checker["scenarios"]:
        item["source_refs"] = [{"document": "requirements.md", "heading_path": ["需求", "UC-001"],
                                "line_start": 5, "line_end": 7}]
    matches = matched(generated, checker, [link("C1", "E1", "full"), link("C1", "D1"),
        link("C2", "E1"), link("C2", "E2", "full"), link("CM", "GM", "full"),
        link("CA", "GA", "full"), link("C3", "GM")])
    return generated, checker, matches


def test_exception_only_dedup_and_generation_contributions():
    generated, checker, matches = fixture_scenes()
    metrics = compare_scenes(generated, checker, matches)
    assert metrics["overall"]["miss_rate"]["rate"] == 0
    exception = metrics["exception_overall"]
    assert exception["checker_count"] == 3
    assert exception["generated_count"] == 4
    assert exception["miss_rate"] == {"numerator": 1, "denominator": 3, "rate": 1 / 3}
    assert exception["existing_completeness"] == {"numerator": 3, "denominator": 4, "rate": .75}
    assert exception["missing_checker_ids"] == ["C3"]
    assert exception["missing_scenarios"][0]["source_refs"][0]["heading_path"] == ["需求", "UC-001"]
    assert exception["partial_only_checker_ids"] == []  # both also have a full link
    explicit = metrics["generation_contributions"]["explicit"]
    derived = metrics["generation_contributions"]["concern_derived"]
    assert explicit["matched_checker_ids"] == ["C1", "C2"]
    assert explicit["matched_checker_count"] == 2
    assert explicit["exclusive_checker_ids"] == ["C2"]
    assert explicit["shared_checker_ids"] == ["C1"]
    assert derived["matched_checker_count"] == 1
    assert derived["exclusive_checker_count"] == 0
    assert derived["shared_checker_count"] == 1
    assert derived["miss_rate"]["rate"] == 2 / 3
    for key in ("common.timeout", "external_service.availability"):
        assert metrics["by_concern_exception"][key]["miss_rate"] == exception["miss_rate"]
    assert metrics["by_group_exception"]["common"]["existing_completeness"] == exception["existing_completeness"]


@pytest.mark.parametrize("count,missed,passed", [(20, 1, False), (21, 1, True), (1, 0, True), (0, 0, False)])
def test_strict_exception_acceptance(count, missed, passed):
    checker = dataset([scene(f"C{i}") for i in range(count)])
    generated = dataset([{**scene("G"), "generation_status": "explicit"}])
    matches = matched(generated, checker, [link(f"C{i}", "G") for i in range(count - missed)])
    metrics = compare_scenes(generated, checker, matches)
    assert metrics["acceptance"]["passed"] is passed
    assert metrics["acceptance"]["threshold"] == .05
    if not count:
        assert metrics["exception_overall"]["miss_rate"]["rate"] is None
    matches["complete"] = False
    with pytest.raises(ValueError, match="matching must be complete"):
        compare_scenes(generated, checker, matches)


def test_success_only_cannot_pass_exception_acceptance():
    generated, checker = dataset([scene("G", kind="main_success")]), dataset([scene("C", kind="main_success")])
    metrics = compare_scenes(generated, checker, matched(generated, checker, [link("C", "G", "full")]))
    assert metrics["overall"]["miss_rate"]["rate"] == 0
    assert metrics["exception_overall"]["generated_count"] == 0
    assert metrics["exception_overall"]["existing_completeness"]["rate"] is None
    assert not metrics["acceptance"]["passed"]


def test_export_columns_numeric_formats_all_scenes_and_diagnostics(tmp_path):
    generated, checker, matches = fixture_scenes()
    generated["scenarios"][1]["concern_keys"] = ["common.timeout"]
    matches["input_hash"] = fingerprint([generated, checker])
    metrics = compare_scenes(generated, checker, matches)
    diagnostic = {"source_refs": checker["scenarios"][2]["source_refs"], "trigger": "query timeout",
                  "concern_key": "common.timeout", "subject_node_id": "NODE-1", "relation_id": "EDGE-1",
                  "candidate_id": "CAND-1", "status": "uncovered", "reason": "未匹配", "suggestion": "补充异常行为"}
    metrics["coverage_diagnostics"] = [diagnostic]
    recommendations = {"backend": "agent", "rerank": False, "items": []}
    artifacts = export_reports(tmp_path, generated, checker, matches, metrics, recommendations)
    wb = load_workbook(artifacts["workbook"])
    try:
        ws = wb["指标"]
        headers = {cell.value: cell.column for cell in ws[1]}
        assert headers["关注点说明"] == headers["类别"] + 1
        rows = {ws.cell(i, headers["类别"]).value: i for i in range(2, ws.max_row + 1)}
        total = 2  # first headline row is exception-only
        assert ws.cell(total, headers["类别"]).value == "exception_overall"
        assert ws.cell(total, headers["已有场景数"]).value == 3
        assert ws.cell(total, headers["漏报分子"]).value == 1
        assert ws.cell(total, headers["漏报分母"]).value == 3
        assert ws.cell(total, headers["漏报率"]).value == pytest.approx(1 / 3)
        assert ws.cell(total, headers["当前已有完整率"]).value == .75
        for i in range(2, ws.max_row + 1):
            assert ws.cell(i, headers["漏报率"]).number_format == "0.00%"
            assert ws.cell(i, headers["当前已有完整率"]).number_format == "0.00%"
            assert ws.cell(i, headers["已有场景数"]).number_format != "0.00%"
        assert "延时是否影响" in ws.cell(rows["common.timeout"], headers["关注点说明"]).value
        assert ws.cell(rows["common"], headers["关注点说明"]).value.startswith("通用关注点：")
        assert ws.cell(total, headers["关注点说明"]).alignment.wrap_text
        assert wb["生成场景"].max_row - 1 == len(generated["scenarios"])
        assert wb["已有场景"].max_row - 1 == len(checker["scenarios"])
        assert wb["生成器漏报"].cell(2, 1).value == "C3"
        assert wb["覆盖诊断"].cell(2, 7).value == "uncovered"
        assert wb["覆盖诊断"].cell(2, 9).value == "补充异常行为"
        discrepancy = wb["分类差异"]
        assert discrepancy.max_row - 1 == len(metrics["classification_discrepancies"])
        rows = list(discrepancy.iter_rows(min_row=2, values_only=True))
        difference = next(r for r in rows if r[0] == "C2" and r[1] == "E2")
        assert "external_service.availability" in difference[3] and "external_service.availability" not in difference[4]
        assert "requirements.md" in difference[9] and difference[7] == checker["scenarios"][1]["trigger"]
        attachment_headers = {cell.value: cell.column for cell in wb["关注点挂载"][1]}
        attachment = [row for row in wb["关注点挂载"].iter_rows(min_row=2) if row[1].value == "D1"]
        assert len(attachment) == 2
        assert all(row[attachment_headers["节点"] - 1].value == "NODE-1" for row in attachment)
    finally:
        wb.close()
    md = (tmp_path / "report.md").read_text(encoding="utf-8")
    assert "# 三 Agent 异常场景补全评估" in md
    assert "异常漏报率：1/3" in md
    assert "全部场景对照漏报率：0/5" in md
    assert "补充异常行为" in md
    assert all(f"### {s['scenario_id']}：" in md for s in generated["scenarios"] + checker["scenarios"])
    exported = json.loads((tmp_path / "metrics.json").read_text(encoding="utf-8"))
    assert exported["coverage_diagnostics"] == [diagnostic]
    assert exported["exception_overall"] == metrics["exception_overall"]


def test_export_legacy_metrics_snapshot_without_new_fields(tmp_path):
    generated, checker, matches = fixture_scenes()
    metrics = compare_scenes(generated, checker, matches)
    for key in ("exception_overall", "generation_contributions", "by_concern_exception", "by_group_exception", "acceptance"):
        del metrics[key]
    original = copy.deepcopy(metrics)
    export_reports(tmp_path, generated, checker, matches, metrics, {"backend": "agent", "rerank": False, "items": []})
    assert metrics == original
    assert (tmp_path / "report.md").read_text(encoding="utf-8").startswith("# 三 Agent 场景评估")
