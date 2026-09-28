import ast
import copy
import json
import os
import re
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook
from tests.test_scene_completion import sample_model
from tools.scene_completion.assembly import _link_predictions_to_scenarios, _merge_equivalent_database_failures, _scenario_catalog, assemble_results
from tools.scene_completion.concerns import plan_concern_matrix, validate_concern_matrix
from tools.scene_completion.exporters import export_workbooks
from tools.scene_completion.graphs import build_use_case_dependency_graph, render_use_case_dependency_svg
from tools.scene_completion.dependency_layout import dependency_layout
from tools.scene_completion.overview import build_system_composition_semantics
from tools.scene_completion.review import _batch_payload, load_ecnu_env_file, review_concerns
from tools.scene_completion.schemas import validate_scene_model
from tools.scene_completion.ssd import generate_ssd_bundle, validate_ssd
from tools.scene_completion.ssd import write_ssd_bundle
from tools.scene_completion.png_renderer import convert_svg_to_png, find_svg_converter
from tools.scene_completion.svg_renderer import render_system_composition_svg


def test_v6_public_ssd_entrypoints_are_defined_once():
    source = Path(__file__).resolve().parents[1] / "tools" / "scene_completion" / "ssd.py"
    tree = ast.parse(source.read_text(encoding="utf-8"))
    for name in ("generate_ssd_bundle", "validate_ssd", "write_ssd_bundle"):
        definitions = [node for node in tree.body if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name == name]
        assert len(definitions) == 1, f"{name} should have one public definition, got {len(definitions)}"


def test_v6_normalization_and_non_data_routing():
    model = sample_model()
    model["version"] = "6"
    sr = model["use_cases"][0]["architecture"]["sr"]
    sr.update({"service_type": "resource_mutation", "classification_status": "inferred", "classification_basis": "本用例 API 执行订单变更", "source_location": "fixture.md:API-ORDER"})
    model["interactions"].append({"interaction_id": "INT-SR-ROUTE", "use_case_id": "UC-001", "from_node": "system", "to_node": "sr-order", "direction": "outgoing", "message": "调用订单变更 API", "abstract_api_id": "API-ORDER", "layer": "SR", "sequence": 8, "source_step_index": 1})
    normalized = validate_scene_model(model)["normalized_model"]
    rr_services = [
        node for node in normalized["system_composition"]["nodes"]
        if node.get("kind") == "abstract_service" and node.get("layer") == "RR"
    ]
    assert len(rr_services) == len(normalized["use_cases"])

    matrix = plan_concern_matrix(model)
    keys = {item["concern_key"] for item in matrix["items"]}
    assert "service.resource_mutation.business_constraint" in keys
    assert "service_relation.call_order" not in keys  # no explicit, evidence-backed SR dependency in this fixture
    assert any(item["status"] == "pending_review" for item in matrix["items"])
    assert not validate_concern_matrix(model, matrix, require_complete=True)["valid"]


def test_v6_semantic_duplicate_messages_are_rejected():
    model = sample_model()
    model["version"] = "6"
    bundle = generate_ssd_bundle(model, "UC-001")
    duplicate = copy.deepcopy(bundle["fused"])
    duplicate["messages"].append(copy.deepcopy(duplicate["messages"][0]))
    report = validate_ssd(duplicate, model)
    assert not report["valid"]
    assert any("duplicate" in error.lower() for error in report["errors"])


def test_png_outputs_are_generated_when_local_converter_exists():
    if not find_svg_converter():
        return
    with tempfile.TemporaryDirectory() as tmp:
        svg = Path(tmp) / "sample.svg"
        png = Path(tmp) / "sample.png"
        svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="20" height="20"><rect width="20" height="20" fill="red"/></svg>', encoding="utf-8")
        result = convert_svg_to_png(svg, png, require=True)
        assert result["status"] == "rendered"
        assert png.exists() and png.read_bytes().startswith(b"\x89PNG")

    model = sample_model()
    model["version"] = "6"
    bundle = generate_ssd_bundle(model, "UC-001")
    with tempfile.TemporaryDirectory() as tmp:
        manifest = write_ssd_bundle(bundle, model, tmp)
        assert all(Path(item["png"]).exists() for item in manifest["artifacts"].values())


def test_v9_actor_frontend_and_external_service_associations_reach_use_cases():
    model = sample_model()
    model["version"] = "6"
    sr_service = next(node for node in model["system_composition"]["nodes"] if node["node_id"] == "sr-order")
    sr_service.update({"name": "OrderService"})
    model["system_composition"]["edges"].append({"edge_id": "EDGE-EXT", "from_node": "sr-order", "to_node": "ext-service", "relation": "sr_external_dependency"})
    normalized = validate_scene_model(model, raise_on_error=True)["normalized_model"]
    associations = [edge for edge in normalized["system_composition"]["edges"] if edge.get("relation") in {"uses_frontend", "serves_use_case", "external_participates_in", "uses_external_service"}]
    mapping = normalized["frontend_mappings"][0]
    assert any(edge["from_node"] == "user" and edge["to_node"] == mapping["frontend_id"] for edge in associations)
    assert any(edge["from_node"] == mapping["frontend_id"] and edge["use_case_id"] == "UC-001" for edge in associations)
    assert any(edge["to_node"] == "ext-service" and edge["use_case_id"] == "UC-001" for edge in associations)
    assert normalized["supported_devices"] == []
    with tempfile.TemporaryDirectory() as tmp:
        svg = render_system_composition_svg(normalized, Path(tmp) / "system.svg").read_text(encoding="utf-8")
    assert 'data-relation="uses_frontend"' in svg
    assert 'data-relation="serves_use_case"' in svg
    assert 'data-relation="external_participates_in"' in svg
    assert "RR" not in svg and "SR" not in svg and "AR" not in svg
    assert 'data-node-id="system"' not in svg


def test_v9_human_frontends_are_isolated_and_fourteen_use_cases_render_in_grid():
    model = sample_model()
    model["version"] = "6"
    model["system_composition"]["nodes"].append({"node_id": "merchant", "name": "商家", "kind": "human_actor"})
    merchant_case = copy.deepcopy(model["use_cases"][0])
    merchant_case.update({"use_case_id": "UC-002", "use_case_name": "商家管理订单", "actors": ["商家"]})
    merchant_case["architecture"]["rr"]["service_id"] = "rr-service-UC-002"
    merchant_case["architecture"]["rr"]["service_name"] = "商家管理订单"
    merchant_case["architecture"]["sr"]["service_id"] = "sr-order-merchant"
    merchant_case["architecture"]["sr"]["design_use_case_id"] = "SRUC-UC-002-ORDER"
    model["system_composition"]["nodes"].append({"node_id": "sr-order-merchant", "name": "MerchantOrderService", "kind": "abstract_service", "layer": "SR", "use_case_id": "UC-002"})
    model["use_cases"].append(merchant_case)
    for index in range(3, 15):
        case = copy.deepcopy(model["use_cases"][0])
        case.update({"use_case_id": f"UC-{index:03d}", "use_case_name": f"用例{index}"})
        case["architecture"]["rr"]["service_id"] = f"rr-service-{index:03d}"
        case["architecture"]["rr"]["service_name"] = f"用例{index}"
        case["architecture"]["sr"]["service_id"] = f"sr-service-{index:03d}"
        case["architecture"]["sr"]["design_use_case_id"] = f"SRUC-{index:03d}"
        model["system_composition"]["nodes"].append({"node_id": f"sr-service-{index:03d}", "name": f"Service{index}", "kind": "abstract_service", "layer": "SR", "use_case_id": f"UC-{index:03d}"})
        model["use_cases"].append(case)
    normalized = validate_scene_model(model, raise_on_error=True)["normalized_model"]
    mappings = {mapping["actor_id"]: mapping["frontend_id"] for mapping in normalized["frontend_mappings"]}
    assert mappings["user"] != mappings["merchant"]
    with tempfile.TemporaryDirectory() as tmp:
        svg = render_system_composition_svg(normalized, Path(tmp) / "overview.svg").read_text(encoding="utf-8")
    assert svg.count("<ellipse") == 14
    assert "RR" not in svg and "SR" not in svg and "AR" not in svg


def test_v9_crud_dependencies_require_an_explicit_evidence_backed_prerequisite():
    model = sample_model()
    model["entities"] = ["Product"]
    model["use_cases"][0]["main_flow"] = [{"step_index": 1, "text": "创建商品"}]
    second = copy.deepcopy(model["use_cases"][0])
    second.update({"use_case_id": "UC-002", "use_case_name": "查询商品", "main_flow": [{"step_index": 1, "text": "读取已发布商品"}]})
    model["use_cases"].append(second)
    model["use_case_entity_operations"] = [
        {"operation_id": "CREATE-PRODUCT", "use_case_id": "UC-001", "entity": "Product", "operation": "C", "source_step_index": 1, "evidence": "写入商品记录", "source_location": "design.md:10"},
        {"operation_id": "READ-PRODUCT", "use_case_id": "UC-002", "entity": "Product", "operation": "R", "source_step_index": 1, "evidence": "查询商品记录", "source_location": "design.md:20"},
    ]
    graph = build_use_case_dependency_graph(model)
    assert graph["edges"] == []  # Shared entity alone is not a dependency.
    model["use_case_entity_operations"][1].update({"depends_on_operations": ["CREATE-PRODUCT"], "dependency_evidence": "查询前必须存在已创建商品", "dependency_source_location": "design.md:21"})
    graph = build_use_case_dependency_graph(model)
    assert len(graph["edges"]) == 1
    edge = graph["edges"][0]
    assert (edge["from_use_case"], edge["to_use_case"], edge["entity"]) == ("UC-001", "UC-002", "Product")
    assert edge["source_locations"] == ["design.md:10", "design.md:20", "design.md:21"]


def test_v10_split_composition_views_separate_actor_participation_and_dependencies(tmp_path):
    model = sample_model()
    model["version"] = "9"
    model["system_composition"]["edges"].append({"edge_id": "EDGE-EXT-V10", "from_node": "sr-order", "to_node": "ext-service", "relation": "sr_external_dependency"})
    normalized = validate_scene_model(model, raise_on_error=True)["normalized_model"]
    participant = render_system_composition_svg(normalized, tmp_path / "system_composition.svg", view="participation").read_text(encoding="utf-8")
    dependency = render_system_composition_svg(normalized, tmp_path / "system_composition_dependencies.svg", view="dependencies").read_text(encoding="utf-8")
    assert 'data-relation="serves_use_case"' in participant
    assert 'data-relation="use_case_dependency"' not in participant
    assert 'data-relation="frontend_to_use_case_area"' in dependency
    assert 'data-relation="serves_use_case"' not in dependency
    assert 'data-relation="external_participates_in"' in participant
    assert 'data-relation="external_participates_in"' in dependency
    actor_lines = participant.count('data-relation="serves_use_case"')
    assert actor_lines > 0
    assert 'data-actor="user"' in participant
    assert 'stroke="#2563EB"' in participant
    assert "RR" not in participant and "SR" not in participant and "AR" not in participant
    assert "RR" not in dependency and "SR" not in dependency and "AR" not in dependency
    p_semantics = build_system_composition_semantics(normalized, "participation")
    d_semantics = build_system_composition_semantics(normalized, "dependencies")
    assert p_semantics["view_id"] == "system_composition"
    assert d_semantics["view_id"] == "system_composition_dependencies"
    assert p_semantics["participation_relations"]["frontend_to_use_case"]
    assert not d_semantics["participation_relations"]["frontend_to_use_case"]
    # The participation chart intentionally uses a compact multi-column grid;
    # direct actor lines may cross ovals by user request. Dependency edges
    # must still avoid unrelated ovals.
    assert '<line data-relation="serves_use_case"' in participant
    _assert_direct_relation_lines_clear_of_unrelated_ellipses(dependency)


def _assert_orthogonal_relation_paths_clear_of_ellipses(svg_text):
    root = ET.fromstring(svg_text)
    ns = {"s": "http://www.w3.org/2000/svg"}
    ellipses = [
        (float(el.attrib["cx"]), float(el.attrib["cy"]), float(el.attrib["rx"]), float(el.attrib["ry"]))
        for el in root.findall(".//s:ellipse", ns)
    ]
    for path in root.findall(".//s:path", ns):
        if not path.attrib.get("data-relation"):
            continue
        tokens = re.findall(r"[MHV]|-?\d+(?:\.\d+)?", path.attrib["d"])
        command = None
        current = None
        points = []
        index = 0
        while index < len(tokens):
            token = tokens[index]
            if token in {"M", "H", "V"}:
                command = token
                index += 1
            if command == "M":
                current = (float(tokens[index]), float(tokens[index+1]))
                index += 2
                points.append(current)
                command = None
            elif command == "H":
                current = (float(tokens[index]), current[1])
                index += 1
                points.append(current)
                command = None
            elif command == "V":
                current = (current[0], float(tokens[index]))
                index += 1
                points.append(current)
                command = None
        for start, end in zip(points, points[1:]):
            for cx, cy, rx, ry in ellipses:
                if start[1] == end[1]:
                    y = start[1]
                    if abs(y-cy) >= ry:
                        continue
                    half = rx * (1 - ((y-cy)/ry)**2) ** 0.5
                    overlap = min(max(start[0], end[0]), cx+half) - max(min(start[0], end[0]), cx-half)
                else:
                    x = start[0]
                    if abs(x-cx) >= rx:
                        continue
                    half = ry * (1 - ((x-cx)/rx)**2) ** 0.5
                    overlap = min(max(start[1], end[1]), cy+half) - max(min(start[1], end[1]), cy-half)
                assert overlap <= 0.01, f"relation path crosses ellipse interior: {path.attrib}"


def _assert_direct_relation_lines_clear_of_unrelated_ellipses(svg_text):
    root = ET.fromstring(svg_text)
    ns = {"s": "http://www.w3.org/2000/svg"}
    ellipses = [
        (el.attrib.get("data-node-id", ""), float(el.attrib["cx"]), float(el.attrib["cy"]), float(el.attrib["rx"]), float(el.attrib["ry"]))
        for el in root.findall(".//s:ellipse", ns)
    ]
    for line in root.findall(".//s:line", ns):
        if not line.attrib.get("data-relation"):
            continue
        x1, y1 = float(line.attrib["x1"]), float(line.attrib["y1"])
        x2, y2 = float(line.attrib["x2"]), float(line.attrib["y2"])
        for _, cx, cy, rx, ry in ellipses:
            if not rx or not ry:
                continue
            # Sample the open segment; endpoint ellipses are allowed only for
            # the intended target/source, never for unrelated use cases.
            hit = False
            for step in range(1, 100):
                t = step / 100
                x, y = x1 + (x2-x1)*t, y1 + (y2-y1)*t
                if ((x-cx)/rx)**2 + ((y-cy)/ry)**2 < 0.995:
                    hit = True
                    break
            if hit:
                start_radius = ((x1-cx)/rx)**2 + ((y1-cy)/ry)**2
                end_radius = ((x2-cx)/rx)**2 + ((y2-cy)/ry)**2
                touches_endpoint = 0.90 <= start_radius <= 1.10 or 0.90 <= end_radius <= 1.10
                assert touches_endpoint, f"direct relation line crosses unrelated ellipse: {line.attrib}"


def test_v10_dependency_layout_marks_cycles_and_renders_cycle_edges(tmp_path):
    nodes = [{"use_case_id": "A"}, {"use_case_id": "B"}, {"use_case_id": "C"}]
    edges = [
        {"edge_id": "E1", "from_use_case": "A", "to_use_case": "B"},
        {"edge_id": "E2", "from_use_case": "B", "to_use_case": "A"},
        {"edge_id": "E3", "from_use_case": "B", "to_use_case": "C"},
    ]
    layout = dependency_layout(nodes, edges)
    assert layout["cycles"] == [["A", "B"]]
    graph = {"project": "fixture", "nodes": nodes, "edges": [dict(edge, cycle_requires_review=edge["edge_id"] in layout["cycle_edge_ids"]) for edge in edges], "layout": layout}
    svg = render_use_case_dependency_svg(graph, tmp_path / "dependencies.svg").read_text(encoding="utf-8")
    assert 'data-cycle="true"' in svg
    assert 'stroke="#C44536"' in svg
    assert '<line data-relation="use_case_dependency"' in svg


def test_v10_test_scenarios_json_covers_catalog_and_preserves_null_predictions(tmp_path):
    model = sample_model()
    model["version"] = "9"
    uc = model["use_cases"][0]
    uc["scenarios"] = [
        {"scenario_id": "main", "scenario_type": "main", "steps": uc["main_flow"]},
        {"scenario_id": "alt", "scenario_type": "alternative", "steps": [{"step_index": 1, "text": "默认查询"}]},
        {"scenario_id": "err", "scenario_type": "requirement_exception", "anchor_step_index": 1, "steps": [{"step_index": 1, "text": "订单不存在"}], "trigger": "订单不存在", "expected_result": "提示不存在", "recovery": "结束", "source_location": "fixture.md:9"},
    ]
    matrix = plan_concern_matrix(model)
    for item in matrix["items"]:
        item.update({"status": "not_applicable", "basis": "无异常证据", "evidence_types": ["ssd"]})
    bundle = assemble_results(model, matrix, {"findings": []})
    artifacts = export_workbooks(bundle, tmp_path)
    exported = json.loads(Path(artifacts["test_scenarios"]).read_text(encoding="utf-8"))
    catalog_ids = {item["scenario_id"] for item in bundle["scenario_catalog"]}
    assert {item["scenario_id"] for item in exported["scenarios"]} == catalog_ids
    assert exported["scenario_count"] == len(bundle["scenario_catalog"])
    main = next(item for item in exported["scenarios"] if item["scenario_type"] == "main_success")
    alternative = next(item for item in exported["scenarios"] if item["scenario_type"] == "alternative")
    assert main["prediction_id"] is None and alternative["prediction_id"] is None
    assert main["steps"] and main["steps"][0]["step_number"] == 1
    assert Path(artifacts["system_composition_dependencies"]).exists()


def test_v7_matrix_workbook_has_real_exchange_uc_and_timeout_only_tab():
    model = sample_model()
    model["version"] = "6"
    matrix = plan_concern_matrix(model)
    for item in matrix["items"]:
        item.update({"status": "not_applicable", "basis": "该交换没有此类异常依据", "evidence_types": ["ssd"]})
    timeout = next(item for item in matrix["items"] if item["concern_key"] == "common.timeout")
    timeout.update({"status": "needs_requirement", "basis": "缺少服务时限要求", "requirement_impact": "unknown", "subsequent_behavior_impact": "unknown", "environment_coordination_impact": "unknown"})
    assert validate_concern_matrix(model, matrix, require_complete=True)["valid"]
    bundle = assemble_results(model, matrix, {"findings": []})
    with tempfile.TemporaryDirectory() as tmp:
        paths = export_workbooks(bundle, tmp)
        workbook = load_workbook(paths["scenario_workbook"], read_only=True, data_only=True)
        assert workbook.sheetnames == ["场景清单", "关注点矩阵", "超时判断"]
        concern = workbook["关注点矩阵"]
        headers = [concern.cell(3, col).value for col in range(1, concern.max_column + 1)]
        assert "交互ID" not in headers
        assert any("用例ID" in str(value) for value in headers)
        assert all(concern.cell(row, headers.index("用例ID") + 1).value for row in range(4, concern.max_row + 1))
        timeout_sheet = workbook["超时判断"]
        timeout_headers = [timeout_sheet.cell(3, col).value for col in range(1, timeout_sheet.max_column + 1)]
        assert "需求满足影响" in timeout_headers
        assert timeout_sheet.max_row - 3 == sum(item["concern_key"] == "common.timeout" for item in matrix["items"])
        impact_col = timeout_headers.index("需求满足影响") + 1
        impact_values = [timeout_sheet.cell(row, impact_col).value for row in range(4, timeout_sheet.max_row + 1)]
        assert "待需求确认" in impact_values
        assert all(value in {"待需求确认", "—", "yes", "no"} for value in impact_values)
        matrix_sheet = workbook["关注点矩阵"]
        matrix_headers = [matrix_sheet.cell(3, col).value for col in range(1, matrix_sheet.max_column + 1)]
        message_col = matrix_headers.index("交互消息/依赖证据") + 1
        assert all(matrix_sheet.cell(row, message_col).value for row in range(4, matrix_sheet.max_row + 1))


def test_v7_source_scenarios_survive_version_six_and_get_exception_predictions():
    model = sample_model()
    model["version"] = "6"
    model["use_cases"][0]["scenarios"] = [
        {"scenario_id": "UC-001-main", "scenario_type": "main", "anchor_step_index": 0, "steps": model["use_cases"][0]["main_flow"]},
        {"scenario_id": "UC-001-alt", "scenario_type": "alternative", "anchor_step_index": 1, "steps": [{"step_index": 1, "text": "默认查询订单"}]},
        {"scenario_id": "UC-001-error", "scenario_type": "requirement_exception", "anchor_step_index": 2, "anchor_label": "2.a", "steps": [{"step_index": 1, "text": "查询订单"}, {"step_index": 2, "text": "订单不存在，系统提示"}], "trigger": "订单不存在", "expected_result": "提示订单不存在", "recovery": "结束分支", "source_location": "fixture.md:10"},
    ]
    matrix = plan_concern_matrix(model)
    for item in matrix["items"]:
        item.update({"status": "not_applicable", "basis": "该交换没有此类异常依据", "evidence_types": ["ssd"]})
    bundle = assemble_results(model, matrix, {"findings": []})
    types = [item["scenario_type"] for item in bundle["scenario_catalog"]]
    assert "main_success" in types
    assert "alternative" in types
    exception = next(item for item in bundle["scenario_catalog"] if item["scenario_type"] == "requirement_exception")
    assert exception["prediction_id"]
    assert exception["concern_key"] == ""
    assert exception["concern"] == "需求异常｜待分类"
    assert any(item["scenario_id"] == exception["scenario_id"] for item in bundle["findings"])


def test_spec_exception_branch_maps_and_deduplicates_concern_findings():
    model = sample_model()
    model["version"] = "8"
    uc = model["use_cases"][0]
    uc["scenarios"] = [
        {"scenario_id": "UC-001-main", "scenario_type": "main", "steps": uc["main_flow"]},
        {"scenario_id": "UC-001-2.a", "scenario_type": "requirement_exception", "anchor_step_index": 2, "anchor_label": "2.a", "steps": [{"step_index": 1, "text": "查询资源"}, {"step_index": 2, "text": "异常触发：资源不存在"}, {"step_index": 2, "text": "返回资源不存在"}], "trigger": "资源不存在", "expected_result": "返回资源不存在", "recovery": "结束", "source_location": "fixture.md:2"},
        {"scenario_id": "UC-001-2.b", "scenario_type": "requirement_exception", "anchor_step_index": 2, "anchor_label": "2.b", "steps": [{"step_index": 1, "text": "查询资源"}, {"step_index": 2, "text": "异常触发：资源已失效"}, {"step_index": 2, "text": "返回资源已失效"}], "trigger": "资源已失效", "expected_result": "返回资源已失效", "recovery": "结束", "source_location": "fixture.md:3"},
    ]
    findings = [
        {"use_case_id": uc["use_case_id"], "source_step_index": 2, "concern_key": "service.query_retrieval.resource_existence", "concern": "资源存在性", "candidate_id": "CAND-A", "exception_type": "ResourceNotFound", "exception_desc": "场景 UC-001-2.a：资源不存在", "trigger": "资源不存在", "expected_result": "返回资源不存在", "basis": "明确引用 UC-001-2.a", "exchange_id": "EX-A", "ssd_message_id": "MSG-A"},
        {"use_case_id": uc["use_case_id"], "source_step_index": 2, "concern_key": "internal_database.resource_existence", "concern": "资源存在性", "candidate_id": "CAND-B", "exception_type": "ResourceNotFound", "exception_desc": "数据库查询未找到资源", "trigger": "查询的资源不存在", "expected_result": "返回资源不存在", "basis": "数据库查询", "exchange_id": "EX-B", "ssd_message_id": "MSG-B"},
        {"use_case_id": uc["use_case_id"], "source_step_index": 2, "concern_key": "service.query_retrieval.resource_existence", "concern": "资源存在性", "candidate_id": "CAND-C", "exception_type": "ResourceUnavailable", "exception_desc": "场景 UC-001-2.b：资源已失效", "trigger": "资源已失效", "expected_result": "返回资源已失效", "basis": "明确引用 UC-001-2.b", "exchange_id": "EX-C", "ssd_message_id": "MSG-C"},
    ]
    catalog = _scenario_catalog(findings, model, None, [{"name": "fixture.md", "text": "1. success\n2. 资源不存在，返回资源不存在\n3. 资源已失效，返回资源已失效"}])
    predictions = _link_predictions_to_scenarios(findings, catalog, model)
    branches = [item for item in catalog if item["scenario_type"] == "requirement_exception"]
    assert len(branches) == 2
    assert all(item["spec_explicitness"] == "yes" for item in branches)
    not_found = next(item for item in branches if item["source_scenario_id"] == "UC-001-2.a")
    unavailable = next(item for item in branches if item["source_scenario_id"] == "UC-001-2.b")
    assert not_found["concern_key"] == "service.query_retrieval.resource_existence"
    assert len(not_found["merged_candidate_ids"]) == 2
    assert unavailable["concern_key"] == "service.query_retrieval.resource_existence"
    assert not_found["prediction_id"] != unavailable["prediction_id"]
    assert len([item for item in predictions if item["scenario_id"] == not_found["scenario_id"]]) == 1
    assert not_found["scenario_source"] == "spec_exception_branch+concern_mapping"


def test_v7_requirement_exception_trace_nodes_are_resolved_from_ssd(tmp_path):
    model = sample_model()
    model["version"] = "6"
    model["use_cases"][0]["scenarios"] = [{
        "scenario_id": "UC-001-main", "scenario_type": "main", "anchor_step_index": 0,
        "steps": model["use_cases"][0]["main_flow"],
    }, {
        "scenario_id": "UC-001-error", "scenario_type": "requirement_exception", "anchor_step_index": 2,
        "anchor_label": "2.a", "steps": [{"step_index": 1, "text": "提交查询"}, {"step_index": 2, "text": "系统拒绝请求"}],
        "trigger": "参数非法", "expected_result": "返回400", "recovery": "修改后重试", "source_location": "fixture.md:10",
    }]
    bundle_ssd = generate_ssd_bundle(model, "UC-001")
    manifest = write_ssd_bundle(bundle_ssd, model, tmp_path / "ssd")
    matrix = plan_concern_matrix(model)
    for item in matrix["items"]:
        item.update({"status": "not_applicable", "basis": "该交换没有此类异常依据", "evidence_types": ["ssd"]})
    assembled = assemble_results(model, matrix, {"findings": []}, {"use_cases": [manifest]})
    exception = next(item for item in assembled["scenario_catalog"] if item["scenario_type"] == "requirement_exception")
    assert exception["source_node"] and exception["target_node"]
    prediction = next(item for item in assembled["findings"] if item.get("exception_origin") == "requirement_branch")
    assert prediction["source_node"] and prediction["target_node"]


def test_v7_database_availability_merges_same_failure_and_preserves_all_trace_refs():
    model = sample_model()
    target = "int-db"
    common = {
        "use_case_id": "UC-001", "source_step_index": 2, "target_node": target,
        "target_node_name": "内部数据库", "concern_key": "internal_database.availability",
        "expected_result": "请求失败并返回服务暂不可用。", "recovery": "提示稍后重试。",
        "trigger": "数据库连接失败。", "source_node": "impl-order", "source_node_name": "订单服务",
        "layer": "AR", "interaction_message": "查询资源", "source_location": "fixture.md:12",
    }
    items = [
        {**common, "exception_type": "数据库连接失败", "exception_desc": "商品查询失败。", "exchange_id": "EX-1", "ssd_message_id": "MSG-1"},
        {**common, "exception_type": "数据库不可用", "exception_desc": "库存读取失败。", "exchange_id": "EX-2", "ssd_message_id": "MSG-2"},
        {**common, "exception_type": "DatabaseUnavailableException", "exception_desc": "价格读取失败。", "exchange_id": "EX-3", "ssd_message_id": "MSG-3"},
    ]
    merged = _merge_equivalent_database_failures(items, validate_scene_model(model)["normalized_model"])
    assert len(merged) == 1
    assert merged[0]["exception_type"] == "数据库不可用"
    assert set(merged[0]["related_exchange_ids"]) == {"EX-1", "EX-2", "EX-3"}
    assert set(merged[0]["related_ssd_message_ids"]) == {"MSG-1", "MSG-2", "MSG-3"}
    different_recovery = {**items[1], "recovery": "由管理员恢复后重新查询。"}
    separate = _merge_equivalent_database_failures([items[0], different_recovery], validate_scene_model(model)["normalized_model"])
    assert len(separate) == 2
    paraphrased = [
        {**items[0], "expected_result": "系统应提示错误信息，如‘商品加载失败’", "recovery": "提示用户稍后重试，恢复后继续主流程。"},
        {**items[1], "expected_result": "系统提示‘商品加载失败，请稍后重试’", "recovery": "提示后结束或回到主流程。"},
        {**items[2], "expected_result": "商品查询失败，系统返回错误信息", "recovery": "按该分支处理后结束或回到主流程。"},
    ]
    merged_paraphrases = _merge_equivalent_database_failures(paraphrased, validate_scene_model(model)["normalized_model"])
    assert len(merged_paraphrases) == 1
    assert set(merged_paraphrases[0]["related_exchange_ids"]) == {"EX-1", "EX-2", "EX-3"}


def test_v7_batch_review_uses_exchange_batches_and_resumes_without_network(tmp_path, monkeypatch):
    model = sample_model()
    model["version"] = "6"
    model["use_cases"][0]["architecture"]["ar"][0]["microservice_id"] = "compute"
    bundle = generate_ssd_bundle(model, "UC-001")
    ssd_root = tmp_path / "ssds"
    manifest = write_ssd_bundle(bundle, model, ssd_root)
    matrix = plan_concern_matrix(model, ssd_manifest={"use_cases": [manifest]})
    for item in matrix["items"]:
        if item["concern_key"] != "common.timeout":
            item.update({"status": "not_applicable", "basis": "本次批测仅验证超时候选", "evidence_types": ["ssd"]})
    monkeypatch.setenv("ECNU_MAX_API_KEY", "not-a-real-secret")
    calls = []

    def fake_post(url, api_key, body, timeout, retries):
        calls.append((url, api_key, body))
        payload = json.loads(body["messages"][1]["content"].split("输入数据：", 1)[1])
        candidate = payload["candidates"][0]
        if len(calls) == 1:
            malformed = {"items": [{"candidate_id": candidate["candidate_id"], "concern_key": candidate["concern_key"], "status": "needs_requirement"}]}
            return {"choices": [{"message": {"content": json.dumps(malformed)}}]}
        response = {"items": [
            {"candidate_id": item["candidate_id"], "concern_key": item["concern_key"], "status": "needs_requirement", "basis": "资料没有给出适用判定的明确依据", "evidence_types": ["requirement"], "findings": [], "requirement_impact": "", "subsequent_behavior_impact": "", "environment_coordination_impact": ""}
            for item in payload["candidates"]
        ]}
        return {"choices": [{"message": {"content": json.dumps(response)}}]}

    config = {"base_url": "https://example.invalid/v1", "model": "mock", "api_key_env": "ECNU_MAX_API_KEY", "max_concurrency": 2, "json_mode": False}
    output = tmp_path / "reviewed.json"
    first = review_concerns(model, {"use_cases": [manifest]}, matrix, config, output, post=fake_post)
    expected_batches = len({item["exchange_id"] for item in matrix["items"] if item["status"] == "pending_review"})
    assert first["valid"] and len(calls) == expected_batches + 1
    assert first["metrics"]["batch_count"] == expected_batches
    assert first["metrics"]["validation_retries"] == 1
    assert first["metrics"]["checkpoint_hits"] == 0
    call_count = len(calls)
    second = review_concerns(model, {"use_cases": [manifest]}, matrix, config, output, post=fake_post)
    assert second["valid"] and len(calls) == call_count
    assert second["metrics"]["checkpoint_hits"] == expected_batches
    assert all(call[1] == "not-a-real-secret" for call in calls)

    assembled_matrix = json.loads(output.read_text(encoding="utf-8"))
    candidate = next(item for item in assembled_matrix["items"] if item["concern_key"] != "common.timeout")
    candidate.update({
        "status": "applicable",
        "basis": "SSD/API evidence confirms a validated input boundary",
        "evidence_types": ["ssd", "api_contract"],
        "findings": [{
            "exception_type": "invalid input",
            "exception_desc": "An invalid request value is rejected.",
            "trigger": "Submit a value outside the interface contract.",
            "scenario_steps": ["Submit invalid data.", "The system returns a validation error."],
            "recovery": "Correct the value and retry.",
            "source_step_index": candidate.get("source_step_index") or 1,
        }],
    })
    assembled = assemble_results(model, assembled_matrix, {"findings": []}, {"use_cases": [manifest]})
    assert assembled["findings"]
    assert assembled["findings"][0]["exchange_id"] == candidate["exchange_id"]


def test_v7_review_payload_includes_matching_api_contract_constraints():
    model = sample_model()
    model["interfaces"] = [{
        "name": "API-S-IF1", "abstract_api_id": "API-S-IF1", "method": "GET", "path": "/products",
        "validation_rules": [{"field": "pageSize", "rule": "1<=pageSize<=100", "source_location": "design.md:12"}],
    }]
    candidate = {"use_case_id": "UC-001", "concern_key": "api.data.range"}
    exchange = {"use_case_id": "UC-001", "messages": [{
        "message_id": "REQ-1", "exchange_id": "EX-1", "message_kind": "request", "abstract_api_id": "API-S-IF1", "message": "查询商品",
    }]}
    payload = _batch_payload(model, "EX-1", [candidate], exchange)
    assert payload["api_contracts"][0]["validation_rules"][0]["field"] == "pageSize"


def test_review_agent_mode_exports_packets_and_merge_validates_agent_results(tmp_path):
    model = sample_model()
    model["version"] = "6"
    bundle = generate_ssd_bundle(model, "UC-001")
    manifest = write_ssd_bundle(bundle, model, tmp_path / "ssds")
    matrix = plan_concern_matrix(model, ssd_manifest={"use_cases": [manifest]})
    for item in matrix["items"]:
        item.update({"status": "not_applicable", "basis": "fixture exclusion", "evidence_types": ["ssd"]})
    candidate = next(item for item in matrix["items"] if item["concern_key"] != "common.timeout")
    candidate.update({"status": "pending_review", "basis": "", "evidence_types": []})

    agent_matrix_path = tmp_path / "agent_matrix.json"
    exported = review_concerns(model, {"use_cases": [manifest]}, matrix, {}, agent_matrix_path, mode="agent")
    assert not exported["valid"]
    packet_dir = Path(exported["agent_batches"]["directory"])
    packet_manifest = json.loads((packet_dir / "manifest.json").read_text(encoding="utf-8"))
    assert packet_manifest["batch_count"] == 1
    packet = json.loads((packet_dir / packet_manifest["batches"][0]["packet"]).read_text(encoding="utf-8"))
    assert packet["input"]["candidates"][0]["concern_key"] == candidate["concern_key"]

    results = {"batches": [{"exchange_id": candidate["exchange_id"], "items": [{
            "candidate_id": candidate["candidate_id"], "concern_key": candidate["concern_key"], "status": "not_applicable",
        "basis": "本交互没有此类异常证据", "evidence_types": ["ssd"], "findings": [],
    }]}]}
    merged = review_concerns(
        model, {"use_cases": [manifest]}, json.loads(agent_matrix_path.read_text(encoding="utf-8")),
        {}, tmp_path / "merged.json", mode="merge-agent", agent_results=results,
    )
    assert merged["valid"]
    merged_matrix = json.loads((tmp_path / "merged.json").read_text(encoding="utf-8"))
    reviewed = next(item for item in merged_matrix["items"] if item["concern_key"] == candidate["concern_key"] and item["exchange_id"] == candidate["exchange_id"])
    assert reviewed["status"] == "not_applicable"


def test_auto_review_falls_back_without_credentials_without_network(tmp_path, monkeypatch):
    monkeypatch.delenv("ECNU_MAX_API_KEY", raising=False)
    model = sample_model()
    model["version"] = "6"
    bundle = generate_ssd_bundle(model, "UC-001")
    manifest = write_ssd_bundle(bundle, model, tmp_path / "ssds")
    matrix = plan_concern_matrix(model, ssd_manifest={"use_cases": [manifest]})
    output = tmp_path / "auto_matrix.json"
    result = review_concerns(model, {"use_cases": [manifest]}, matrix, {}, output, mode="auto")
    assert not result["valid"]
    assert result["fallback_reason"] == "API key environment variable is not set: ECNU_MAX_API_KEY"
    assert result["agent_batches"]["batch_count"] > 0
    saved = json.loads(output.read_text(encoding="utf-8"))
    assert all(item["status"] == "pending_review" for item in saved["items"])


def test_auto_review_falls_back_only_failed_external_exchanges(tmp_path, monkeypatch):
    monkeypatch.setenv("ECNU_MAX_API_KEY", "test-placeholder")
    model = sample_model()
    model["version"] = "6"
    bundle = generate_ssd_bundle(model, "UC-001")
    manifest = write_ssd_bundle(bundle, model, tmp_path / "ssds")
    matrix = plan_concern_matrix(model, ssd_manifest={"use_cases": [manifest]})
    expected_batches = len({item["exchange_id"] for item in matrix["items"] if item["status"] == "pending_review"})
    calls = []

    def failed_post(url, api_key, body, timeout, retries):
        calls.append(url)
        raise RuntimeError("simulated network outage")

    result = review_concerns(
        model, {"use_cases": [manifest]}, matrix,
        {"base_url": "https://example.invalid/v1", "model": "mock", "max_concurrency": 2},
        tmp_path / "auto_failed.json", mode="auto", post=failed_post,
    )
    assert len(calls) == expected_batches
    assert not result["valid"]
    assert result["agent_batches"]["batch_count"] == expected_batches
    saved = json.loads((tmp_path / "auto_failed.json").read_text(encoding="utf-8"))
    assert all(item["status"] == "pending_review" for item in saved["items"])

def test_ecnu_env_file_loader_only_reads_standard_keys(tmp_path, monkeypatch):
    for key in ("ECNU_MAX_MODEL", "ECNU_MAX_API_KEY", "ECNU_MAX_BASE_URL"):
        monkeypatch.delenv(key, raising=False)
    env_file = tmp_path / ".env"
    env_file.write_text(
        "ECNU_MAX_MODEL=mock-model\nECNU_MAX_API_KEY=not-a-real-secret\n"
        "ECNU_MAX_BASE_URL=https://example.invalid/v1\nUNRELATED=ignored\n",
        encoding="utf-8",
    )
    loaded = load_ecnu_env_file(env_file)
    assert set(loaded) == {"ECNU_MAX_MODEL", "ECNU_MAX_API_KEY", "ECNU_MAX_BASE_URL"}
    assert os.environ["ECNU_MAX_MODEL"] == "mock-model"
    assert "UNRELATED" not in loaded
