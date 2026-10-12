from copy import deepcopy
from types import SimpleNamespace
import json
import pytest
from openpyxl import load_workbook
from scene_completion.concern_gap import PROTOCOL,locations,merge_units,metrics
from scene_completion.concern_gap_ranking import rank
from scene_completion.concern_gap_replay import verify_bundle,rebuild_snapshot
from scene_completion.interaction_protocol import build_ssd,route_mounts
from scene_completion.sources import index_sources,fingerprint
from scene_completion.semantic_backend import write_json
from tests.test_interaction_protocol import model,reference
from tests.test_concern_gap import unit,Vectors


def bundle(tmp_path):
    src=tmp_path/'rr.md';src.write_text('# Requirements\n## UC-1\n用户提交商品编号。商品编号必须存在。\n',encoding='utf8')
    index=index_sources([str(src)]);index['documents'][0]['requirement_layer']='RR'
    m=model(index);ssd=build_ssd(m);mounts=route_mounts(ssd)
    by_exchange={e['exchange_id']:e for e in ssd['exchanges']}
    loc={eid:r['location_id'] for r in locations(ssd) for eid in r['exchange_ids']}
    source=[];units=[];checks=[]
    for i,mount in enumerate(mounts):
        mount['mount_reason']='normal operation';mount['feature_evidence']=[]
        s=unit('g'+str(i),cause=mount['concern_key'].replace('.','_'),location=loc[mount['exchange_id']],keys=[mount['concern_key']])
        s.update(use_case_id='UC-1',requirement_layer=mount['requirement_layer'],source_refs=reference(index),
            generation_source='ssd_concern',
            exchange_id=mount['exchange_id'],mount_id=mount['mount_id'],message_ids=mount['message_ids'],
            subject_node_id=mount['subject_node_id'],payload_direction=mount['payload_direction'])
        source.append(s)
        u=deepcopy(s);u.update(source_scenario_id=s['scenario_id'],unit_id='u'+str(i),concern_key=mount['concern_key'],reason='source',evidence_quote=s['trigger_condition'])
        units.append(u);checks.append({'unit_id':u['unit_id'],'passed':True})
    g=merge_units(units,'G2');ranking=rank(g,[],Vectors())
    rel={'EG':[],'TG':[],'TE':[]};result=metrics(g,[],[],ranking['selected_ids'],rel)
    out={'protocol':PROTOCOL,'scope':{'kind':'representative','use_case_ids':['UC-1'],'granularity':'complete_use_cases'},
        'source_index':index,'test_source_index':index,'model':m,'ssd':ssd,'mounts':mounts,'features':[],
        'graphs':[],'normal':[],'G':g,'E':[],'history':[],'T':[],'ranking':ranking,'relations':rel,'metrics':result,
        'source_generated':source,'source_existing':[],'source_tests':[],'source_units':{'G':units,'E':[],'T':[]},
        'reverse_checks':{'G':{'reviews':checks},'E':{'reviews':[]},'T':{'reviews':[]}}}
    out['result_hash']=fingerprint(out)
    return out


def test_offline_rebuild_validates_full_mounts_and_exports_chinese_headers(tmp_path):
    value=bundle(tmp_path);assert verify_bundle(value)['complete'] is True
    path=tmp_path/'snapshot.json';write_json(path,{'payload':value,'payload_hash':fingerprint(value)})
    result=rebuild_snapshot(path,tmp_path/'rebuilt')
    assert result['result_hash']==value['result_hash'] and result['execution_evidence_verified'] is False
    book=load_workbook(tmp_path/'rebuilt/scene_assessment.xlsx',read_only=True,data_only=True)
    try:
        headers=next(book['指标'].iter_rows(values_only=True))
        assert headers[:2]==('类别','关注点说明')
        assert '缺失异常' in book.sheetnames and '最终推荐' in book.sheetnames and '测试GT' in book.sheetnames
        assert next(book['指标'].iter_rows(min_row=2,values_only=True))[6] is None
    finally:book.close()


def test_snapshot_hash_and_missing_mounts_rejected_before_export(tmp_path):
    original=bundle(tmp_path)
    tampered=deepcopy(original);tampered['mounts'].pop()
    with pytest.raises(ValueError,match='hash changed'):verify_bundle(tampered)
    tampered['result_hash']=fingerprint({k:v for k,v in tampered.items() if k!='result_hash'})
    with pytest.raises(ValueError,match='mounts'):verify_bundle(tampered)


def test_seed_or_rank_tampering_is_not_accepted_as_new_results(tmp_path):
    value=bundle(tmp_path);value['ranking']['selected_ids']=[]
    value['result_hash']=fingerprint({k:v for k,v in value.items() if k!='result_hash'})
    with pytest.raises(ValueError,match='TopN'):verify_bundle(value)


def test_rejected_reverse_record_cannot_be_merged_into_published_set(tmp_path):
    value=bundle(tmp_path);value['reverse_checks']['G']['reviews'][0]['passed']=False
    value['result_hash']=fingerprint({k:v for k,v in value.items() if k!='result_hash'})
    with pytest.raises(ValueError,match='reverse checks'):verify_bundle(value)


def test_replay_keeps_all_batch_identities_when_inverse_repair_ledgers_are_empty(tmp_path,monkeypatch):
    from scene_completion import concern_gap_replay as replay_module, interaction_cli
    from scene_completion import concern_gap_batches as b
    from scene_completion.semantic_backend import prepare_packets,run_agent_packets
    value=bundle(tmp_path)
    value['normalization_repairs']={c:[] for c in ('G','E','T')}
    value['base_source_units']=deepcopy(value['source_units'])
    value['result_hash']=fingerprint({k:v for k,v in value.items() if k!='result_hash'})
    write_json(tmp_path/'result_bundle.json',value)
    write_json(tmp_path/'run_manifest.json',{'complete':True,'protocol':PROTOCOL,
        'versions':b.versions(),'result_hash':value['result_hash'],'scope':value['scope']})
    write_json(tmp_path/'sources_index.json',value['source_index'])
    write_json(tmp_path/'generator_base_identity.json',{'use_cases':[{'use_case_id':'UC-1'}]})
    write_json(tmp_path/'historical_ssd.json',value['ssd'])
    outputs={'interaction-model':value['model']['use_cases'], 'gap-generation':value['source_generated'],
        'gap-features':value['features'],'gap-graph-check':value['graphs'],
        'interaction-participants':[], 'interaction-checker':value['source_existing'],
        'interaction-checker-audit':[], 'gap-test-extract':value['source_tests']}
    for c in ('G','E','T'):
        outputs['gap-normalize-'+c]=value['source_units'][c]
        outputs['gap-reverse-check-'+c]=value['reverse_checks'][c]['reviews']
    for name in ('EG','TG','TE'):outputs['gap-ambiguity-'+name]=[]
    specs={}
    for name,items in outputs.items():
        def prepare(directory,n=name):return prepare_packets(directory,n,[{'fixture_stage':n}],[],{'items':[]})
        validator=lambda p,r:r['items']
        folder=tmp_path/'batches'/name
        prepare(folder)
        class FixtureClient:
            config={'max_concurrency':1,'semantic_max_attempts':1}
            def model(self,kind):return 'ecnu-max'
            def agent(self,p):return {'items':items,'execution_model_evidence':{
                'requested_model':'ecnu-max','thinking':False,'reasoning_effort':None}}
        assert run_agent_packets(folder,FixtureClient(),validator)['complete']
        specs[name]=(prepare,validator)
    monkeypatch.setattr(replay_module,'stage_specs',lambda root:specs)
    monkeypatch.setattr(interaction_cli,'merge_model',lambda *args:value['model'])
    result=replay_module.replay(tmp_path,cache_resume=True)
    assert len(result['batches'])==len(specs)
    assert {r['stage'] for r in result['batches']}==set(specs)
    assert all(r['status']=='cached' for r in result['batches'])
