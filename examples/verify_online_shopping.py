"""Offline verification of the requirement-only online shopping acceptance run."""
from __future__ import annotations

import argparse
import math
import shutil
import sys
import tempfile
from pathlib import Path, PureWindowsPath
from types import SimpleNamespace
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / ".cac" / "tools"))

from scene_completion.comparison import compare_scenes
from scene_completion.coverage_diagnostics import coverage_diagnostics
from scene_completion.generator import generate_scenes
from scene_completion.generator_model import merge_generator_model
from scene_completion.placements import export_placements, annotate_composition_evidence
from scene_completion.overview import build_system_composition_semantics
from scene_completion.matching_audit import validate_native_review_binding, exclude_native_rejections
from scene_completion.role_cli import run_evaluate, validate_execution_models
from scene_completion.role_reports import export_reports
from scene_completion.semantic_backend import read_json, write_json
from scene_completion.sources import fingerprint
from scene_completion.three_roles import merge_matches
from verify_demo import compact


def pack_stage(directory):
    root = Path(directory)
    manifest = read_json(root / "manifest.json")
    return {"manifest": manifest, "packets": {
        b["batch_id"]: read_json(root / "packets" / (b["batch_id"] + ".json")) for b in manifest["batches"]},
        "results": {b["batch_id"]: read_json(root / "results" / (b["batch_id"] + ".json"))
                    for b in manifest["batches"]}}


def unpack_stage(value, directory):
    root = Path(directory)
    ids = {b["batch_id"] for b in value["manifest"]["batches"]}
    if set(value["packets"]) != ids or set(value["results"]) != ids:
        raise ValueError("archived stage has an incomplete packet/result universe")
    write_json(root / "manifest.json", value["manifest"])
    for name in ("packets", "results"):
        for bid, item in value[name].items():
            write_json(root / name / (bid + ".json"), item)


def archived_artifact_loader(root, manifest):
    """Resolve recorded paths exclusively inside the portable snapshot.

    Original paths remain provenance data; changing them would invalidate G
    and its source-bound batch inputs. No access to the original run is needed.
    """
    paths = {}
    for entry in manifest.get("use_cases", []):
        original = ((entry.get("artifacts") or {}).get("fused") or {}).get("json")
        if original:
            name = PureWindowsPath(original).name
            paths[original] = Path(root) / "diagrams" / entry["use_case_id"] / name
    def load(original):
        if original not in paths:
            raise ValueError("unarchived SSD artifact requested")
        return read_json(paths[original])
    return load


def validate_replay(root):
    root = Path(root)
    proof = read_json(root / "replay_verification.json")
    expected = {name: fingerprint(read_json(root / (name + ".json"))) for name in
                ("scene_model", "generated_scenarios", "existing_scenarios", "scenario_matches", "metrics", "recommendations")}
    if (proof.get("complete") is not True or proof.get("same_version") is not True or
            proof.get("network_requests") != 0 or proof.get("before") != expected or proof.get("after") != expected or
            proof.get("tool_versions") != read_json(root / "run_manifest.json").get("tool_versions") or
            proof.get("evaluation") != read_json(root / "evaluation.json")):
        raise ValueError("same-version complete replay differs from archived acceptance inputs")
    expected_stages = {"checker", "generator-taxonomy", "checker-taxonomy", "matching", "recommendation"}
    stages = proof.get("cached_stages", {})
    if set(stages) != expected_stages or proof.get("generator_preparation_batches") != 14:
        raise ValueError("same-version replay omitted required stages")
    for stage, state in stages.items():
        manifest = read_json(root / "batches" / stage / "manifest.json")
        if state.get("complete") is not True or state.get("cached_batches") != len(manifest["batches"]):
            raise ValueError("same-version replay batch universe is incomplete")
    return proof


def validate_report_artifacts(root, regenerated):
    root, regenerated = Path(root), Path(regenerated)
    if (root / "report.md").read_text(encoding="utf-8") != (regenerated / "report.md").read_text(encoding="utf-8"):
        raise ValueError("Markdown report differs from recomputed metrics and recommendations")
    expected = load_workbook(regenerated / "scene_assessment.xlsx", read_only=False, data_only=False)
    actual = load_workbook(root / "scene_assessment.xlsx", read_only=False, data_only=False)
    try:
        if actual.sheetnames != expected.sheetnames:
            raise ValueError("workbook sheets differ from the full export")
        for name in expected.sheetnames:
            left, right = actual[name], expected[name]
            if (left.max_row, left.max_column, left.freeze_panes, left.auto_filter.ref) != (
                    right.max_row, right.max_column, right.freeze_panes, right.auto_filter.ref):
                raise ValueError("workbook layout or row universe differs: " + name)
            for lr, rr in zip(left.iter_rows(), right.iter_rows()):
                for a, b in zip(lr, rr):
                    if (a.value, a.data_type, a.number_format, a.alignment.wrap_text) != (
                            b.value, b.data_type, b.number_format, b.alignment.wrap_text):
                        raise ValueError(f"workbook cell differs: {name}!{a.coordinate}")
            if ({k: v.width for k, v in left.column_dimensions.items()} !=
                    {k: v.width for k, v in right.column_dimensions.items()} or
                    {k: v.height for k, v in left.row_dimensions.items()} !=
                    {k: v.height for k, v in right.row_dimensions.items()}):
                raise ValueError("workbook readable widths/heights differ: " + name)
    finally:
        actual.close()
        expected.close()


def verify(root, output_dir=None):
    root = Path(root)
    with tempfile.TemporaryDirectory(prefix="online-scene-verification-") as directory:
        work = Path(directory)
        for file in ("scene_model.json", "generated_scenarios.json", "existing_scenarios.json",
                     "scenario_matches.json", "recommendations.json", "run_manifest.json",
                     "sources_index.json", "checker_sources_index.json"):
            write_json(work / file, read_json(root / file))
        evidence = read_json(root / "batch_evidence.json")
        write_json(work / "diagrams" / "diagram_manifest.json", read_json(root / "diagram_manifest.json"))
        for stage, packed in evidence["pipeline"].items():
            unpack_stage(packed, work / "batches" / stage)
        native_reviews = read_json(root / "native_reviews.json")
        required_reviews = {"matching", "baseline-matching"}
        if "matching-recovery" in evidence["pipeline"]:
            required_reviews.add("matching-recovery")
        if "baseline-recovery" in evidence:
            required_reviews.add("baseline-recovery")
        if {r["stage"] for r in native_reviews} != required_reviews or len(native_reviews) != len(required_reviews):
            raise ValueError("missing or duplicate final/baseline/recovery native review")
        for entry in native_reviews:
            stage = entry["stage"]
            packed = evidence[stage] if stage.startswith("baseline-") else evidence["pipeline"][stage]
            directory = work / "native-inputs" / stage
            unpack_stage(packed, directory)
            packet = validate_native_review_binding(directory, entry["review"])
            if packet != entry["packet"]:
                raise ValueError("native review packet differs from exact applied proof/input")
            cursor = 0
            for segment in entry["review"].get("review_segments", []):
                result = read_json(root / Path(segment["file"]).name)
                if (segment["start"] != cursor or segment["end"] <= cursor or
                        fingerprint(result) != segment["result_hash"] or
                        result.get("input_hash") != entry["review"]["input_hash"] or
                        result.get("model") != "gpt-6-luna" or result.get("reasoning_effort") != "max" or
                        result["decisions"] != entry["review"]["decisions"][cursor:segment["end"]]):
                    raise ValueError("native review segment differs from the combined decisions")
                cursor = segment["end"]
            if cursor and cursor != len(entry["review"]["decisions"]):
                raise ValueError("native review segments are incomplete")
            if not stage.startswith("baseline-"):
                write_json(work / "batches" / stage / "native_review.json", entry["review"])
        unpack_stage(evidence["generator-model"], work / "prepared" / "batches" / "generator-model")
        seed = read_json(root / "generator_seed.json")
        validate_execution_models(work / "prepared" / "batches" / "generator-model")
        sources = read_json(work / "sources_index.json")
        model = merge_generator_model(seed, sources, work / "prepared" / "batches" / "generator-model")
        if not model.get("complete") or model != read_json(work / "scene_model.json"):
            raise ValueError("archived generator model differs from independent source-bound batches")
        preparation = read_json(root / "preparation_manifest.json")
        if (not preparation.get("complete") or preparation.get("agent_model") != "ecnu-max" or
                preparation.get("source_hash") != fingerprint(sources) or
                preparation.get("skill_hash") != model["generator_provenance"]["skill_hash"]):
            raise ValueError("generator execution provenance differs from the prepared model")
        diagram_manifest = read_json(root / "diagram_manifest.json")
        loader = archived_artifact_loader(root, diagram_manifest)
        regenerated = generate_scenes(model, sources, diagram_manifest, "SR", artifact_loader=loader)
        if "generator-taxonomy" in evidence["pipeline"]:
            from scene_completion.checker_taxonomy import apply_taxonomy_audit
            regenerated = apply_taxonomy_audit(regenerated, sources, work / "batches" / "generator-taxonomy", "generator")
        generated, checker, matches = [read_json(work / (name + ".json")) for name in
                                      ("generated_scenarios", "existing_scenarios", "scenario_matches")]
        if regenerated != generated:
            raise ValueError("archived G differs from unfiltered structured generation")
        architecture = export_placements(model, generated, work,
                                         evidence_review=read_json(root / "architecture_evidence_review.json"))
        if (not architecture["evidence_review_complete"] or architecture != read_json(root / "architecture_changes.json") or
                (work / "architecture_calls.svg").read_text(encoding="utf-8") !=
                (root / "architecture_calls.svg").read_text(encoding="utf-8")):
            raise ValueError("architecture report differs from source-bound evidence review")
        if annotate_composition_evidence(build_system_composition_semantics(model), architecture) != read_json(root / "system_composition.json"):
            raise ValueError("composition source confidence differs from reviewed relationships")
        metrics = compare_scenes(generated, checker, matches)
        metrics["coverage_diagnostics"] = coverage_diagnostics(generated, checker, matches)
        if compact(metrics) != read_json(root / "metrics_summary.json"):
            raise ValueError("metric snapshot differs from recomputed all-scene and exception metrics")
        write_json(work / "metrics.json", metrics)
        # The large full metrics object is rebuilt, not duplicated in the snapshot.
        write_json(work / "replay_verification.json", read_json(root / "replay_verification.json"))
        evaluation, code = run_evaluate(SimpleNamespace(output_dir=str(work), artifact_loader=loader))
        if code or evaluation != read_json(root / "evaluation.json"):
            raise ValueError("requirement-only ecnu-max acceptance gate failed")
        validate_replay(work)
        baseline = root / "baseline"
        old_g, fixed_c = [read_json(baseline / (name + ".json")) for name in
                          ("generated_scenarios", "existing_scenarios")]
        current_c = [s for s in checker["scenarios"] if s["scenario_type"] == "requirement_exception"]
        if len(current_c) != 37 or fixed_c["scenarios"] != current_c or fixed_c["source_hash"] != checker["source_hash"]:
            raise ValueError("before/after comparison does not use the same frozen C")
        unpack_stage(evidence["baseline-matching"], work / "baseline-matching")
        old_matches = merge_matches(old_g, fixed_c, work / "baseline-matching")
        validate_execution_models(work / "baseline-matching")
        if "baseline-recovery" in evidence:
            subset = read_json(baseline / "recovery_checker.json")
            if any(s not in fixed_c["scenarios"] for s in subset["scenarios"]):
                raise ValueError("baseline recovery changed the frozen C")
            unpack_stage(evidence["baseline-recovery"], work / "baseline-recovery")
            validate_execution_models(work / "baseline-recovery")
            recovered = merge_matches(old_g, subset, work / "baseline-recovery")
            if not recovered["complete"]:
                raise ValueError("baseline recovery comparison is incomplete")
            links = {(m["checker_scenario_id"], m["generated_scenario_id"]): m for m in old_matches["matches"]}
            links.update({(m["checker_scenario_id"], m["generated_scenario_id"]): m for m in
                          exclude_native_rejections(recovered["matches"], work / "baseline-matching")})
            old_matches["matches"] = list(links.values())
        if not old_matches["complete"] or old_matches != read_json(baseline / "scenario_matches.json"):
            raise ValueError("baseline matching evidence is incomplete or differs")
        old_metrics = compare_scenes(old_g, fixed_c, old_matches)
        if compact(old_metrics) != read_json(baseline / "metrics_summary.json"):
            raise ValueError("baseline metric snapshot differs")
        rec = read_json(work / "recommendations.json")
        for item in rec["items"]:
            if not math.isclose(item["confidence"], .7 * item["support_score"] + .3 * item["missing_score"]):
                raise ValueError("recommendation score formula differs")
        export_reports(work, generated, checker, matches, metrics, rec)
        validate_report_artifacts(root, work)
        if output_dir:
            Path(output_dir).mkdir(parents=True, exist_ok=True)
            for filename in ("architecture_changes.json", "architecture_evidence_review.json",
                             "architecture_calls.svg", "concern_placements.md", "system_composition.json", "system_composition.svg"):
                shutil.copy2(root / filename, Path(output_dir) / filename)
            export_reports(Path(output_dir), generated, checker, matches, metrics, rec)
        print(f"Verified: G={len(generated['scenarios'])}, requirement exceptions={len(current_c)}, "
              f"miss={metrics['exception_overall']['miss_rate']}, recommendations={len(rec['items'])}")
        print(f"Baseline with the SAME C: miss={old_metrics['exception_overall']['miss_rate']}")
        return evaluation


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--snapshot-dir", default=str(Path(__file__).parent / "online_shopping_demo"))
    parser.add_argument("--output-dir", help="optionally regenerate complete JSON, Markdown and Excel")
    args = parser.parse_args()
    verify(args.snapshot_dir, args.output_dir)


if __name__ == "__main__":
    main()
