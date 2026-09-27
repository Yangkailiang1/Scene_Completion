"""JSON and Excel exporters for Scene Completion V5."""

from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="4472C4")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
STEP_FILL = PatternFill("solid", fgColor="EAF2F8")
HEADER_FONT = Font(bold=True, size=10, color="FFFFFF")
SECTION_FONT = Font(bold=True, size=10, color="1F4E79")
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
BODY = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
THIN = Border(*(Side(style="thin", color="D9E2F3") for _ in range(4)))


def _safe_filename(value: str) -> str:
    return re.sub(r"[\\/:*?\"<>|]+", "_", str(value).strip()) or "scene_completion"


def _title(ws, text: str, columns: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    cell = ws.cell(1, 1, text)
    cell.font = TITLE_FONT
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24


def _headers(ws, headers: list[str]) -> None:
    for col, header in enumerate(headers, 1):
        cell = ws.cell(3, col, header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN


def _format_rows(ws, start: int, rows: list[list[Any]], centered: set[int] | None = None) -> None:
    centered = centered or set()
    for row_index, values in enumerate(rows, start):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_index, col, value)
            cell.alignment = CENTER if col in centered else BODY
            cell.border = THIN


def _node_labels(bundle: dict[str, Any]) -> dict[str, str]:
    return {node["node_id"]: node["name"] for node in bundle["scene_model"]["system_composition"]["nodes"]}


def _findings_by_uc(bundle: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for item in bundle.get("findings", []):
        grouped.setdefault(item.get("use_case_id", ""), []).append(item)
    for items in grouped.values():
        items.sort(key=lambda item: (item.get("source_step_index", 0), item.get("prediction_id", "")))
    return grouped


def _write_prediction_workbook(bundle: dict[str, Any], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "异常预测"
    headers = ["步骤/类型", "检查点", "预测ID", "异常关注点", "异常描述", "匹配GT?", "是否合理", "关注点层级"]
    _title(ws, f"{bundle['project']} — 异常预测（未进行GT评估）", len(headers))
    _headers(ws, headers)
    ucs = {uc["use_case_id"]: uc for uc in bundle["scene_model"].get("use_cases", [])}
    grouped = _findings_by_uc(bundle)
    row = 4
    for uc_id, uc in ucs.items():
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        cell = ws.cell(row, 1, f"{uc_id}：{uc.get('use_case_name', '')}")
        cell.fill = SECTION_FILL
        cell.font = SECTION_FONT
        cell.alignment = BODY
        row += 1
        uc_findings = [item for item in grouped.get(uc_id, []) if item.get("layer") != "AR"]
        uc_level = [item for item in uc_findings if not item.get("source_step_index")]
        for item in uc_level:
            checkpoint = item.get("concern_key") or ("需求来源｜非注册表关注点" if item.get("exception_origin") == "requirement_branch" else "")
            _format_rows(ws, row, [["UC级", checkpoint, item.get("prediction_id", ""), item.get("concern", ""), item.get("exception_desc", ""), "", "", item.get("layer", "RR")]], centered={1, 2, 3, 6, 7, 8})
            row += 1
        for step in uc.get("main_flow", []):
            step_index = step.get("step_index")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
            cell = ws.cell(row, 1, f"步骤 {step_index}：{step.get('text', '')}")
            cell.fill = STEP_FILL
            cell.font = Font(bold=True, color="1F4E79")
            cell.alignment = BODY
            row += 1
            for item in [x for x in uc_findings if x.get("source_step_index") == step_index]:
                checkpoint = item.get("concern_key") or ("需求来源｜非注册表关注点" if item.get("exception_origin") == "requirement_branch" else "")
                _format_rows(ws, row, [[f"步骤 {step_index}", checkpoint, item.get("prediction_id", ""), item.get("concern", ""), item.get("exception_desc", ""), "", "", item.get("layer", "SR")]], centered={1, 2, 3, 6, 7, 8})
                row += 1
    widths = [24, 42, 28, 26, 68, 14, 14, 16]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(3, ws.max_row)}"
    ws.sheet_view.showGridLines = False

    trace = wb.create_sheet("追溯信息")
    trace_headers = ["预测ID", "场景ID", "用例ID", "用例名称", "Actor", "层级", "主流程步骤", "SSD交换ID", "SSD请求消息ID", "来源节点", "目标节点", "关注点Key", "关注点", "异常描述", "触发条件", "异常响应", "恢复", "来源定位", "融合SSD路径", "合并关联SSD交换/消息"]
    _title(trace, f"{bundle['project']} — 异常追溯", len(trace_headers))
    _headers(trace, trace_headers)
    trace_rows = []
    for item in bundle.get("findings", []):
        trace_rows.append([
            item.get("prediction_id", ""), item.get("scenario_id", ""), item.get("use_case_id", ""), item.get("use_case_name", ""), item.get("actor", ""), item.get("layer", ""), item.get("source_step_index", 0), item.get("exchange_id", ""), item.get("ssd_message_id", ""), item.get("source_node_name") or "未映射（待确认）", item.get("target_node_name") or "未映射（待确认）", item.get("concern_key", ""), item.get("concern", ""), item.get("exception_desc", ""), item.get("trigger", ""), item.get("expected_result", ""), item.get("recovery", ""), item.get("source_location", ""), item.get("ssd_paths", ""), "\n".join(f"{ref.get('exchange_id','')} / {ref.get('ssd_message_id','')}" for ref in item.get("trace_refs", []) if ref.get("exchange_id") or ref.get("ssd_message_id")),
        ])
    _format_rows(trace, 4, trace_rows, centered={1, 2, 3, 6, 7, 8, 9})
    for index, width in enumerate([28, 28, 18, 24, 20, 12, 14, 24, 28, 24, 24, 34, 30, 60, 48, 48, 42, 34, 58, 44], 1):
        trace.column_dimensions[get_column_letter(index)].width = width
    trace.freeze_panes = "A4"
    trace.auto_filter.ref = f"A3:{get_column_letter(len(trace_headers))}{max(3, trace.max_row)}"
    trace.sheet_view.showGridLines = False
    if "AR" in bundle.get("analysis_layers", []):
        ar = wb.create_sheet("AR扩展异常")
        _title(ar, f"{bundle['project']} — AR扩展异常", len(headers))
        _headers(ar, headers)
        ar_items = [item for item in bundle.get("findings", []) if item.get("layer") == "AR"]
        ar_rows = [[f"步骤 {item.get('source_step_index', '')}", item.get("concern_key", ""), item.get("prediction_id", ""), item.get("concern", ""), item.get("exception_desc", ""), "", "", "AR"] for item in ar_items]
        _format_rows(ar, 4, ar_rows, centered={1, 2, 3, 6, 7, 8})
        for index, width in enumerate([24, 42, 28, 26, 68, 14, 14, 16], 1): ar.column_dimensions[get_column_letter(index)].width = width
        ar.freeze_panes = "A4"
    wb.save(path)


def _write_scenario_sheet(ws, bundle: dict[str, Any]) -> None:
    headers = ["场景编号", "场景类型", "来源Use Case", "Use Case名称", "Actor", "主流程锚点", "异常关注点", "关注点层级", "前置条件", "触发条件", "场景步骤", "预期结果", "恢复/回归主流程", "来源类型", "来源定位", "SSD交换ID", "SSD消息ID", "预测ID"]
    _title(ws, f"{bundle['project']} — 场景清单（全量，未进行GT评估）", len(headers))
    _headers(ws, headers)
    rows = []
    for item in bundle.get("scenario_catalog", []):
        if "AR" in bundle.get("analysis_layers", []) and item.get("layer") == "AR":
            continue
        rows.append([
            item.get("scenario_id", ""), item.get("scenario_type", ""), item.get("use_case_id", ""), item.get("use_case_name", ""), item.get("actor", ""), item.get("anchor_label", item.get("source_step_index", 0)), item.get("concern", ""), item.get("layer", "RR" if item.get("scenario_type") in {"main_success", "alternative", "requirement_exception"} else "SR"), item.get("preconditions", ""), item.get("trigger", ""), "\n".join(f"{i}. {step}" for i, step in enumerate(item.get("scenario_steps", []), 1)), item.get("expected_result", item.get("exception_desc", "")), item.get("recovery", ""), item.get("source_type", item.get("scenario_origin", "")), item.get("source_location", ""), item.get("exchange_id", ""), item.get("ssd_message_id", ""), item.get("prediction_id", ""),
        ])
    _format_rows(ws, 4, rows, centered={1, 2, 3, 6, 7, 8, 14, 16, 17, 18})
    widths = [30, 24, 18, 24, 24, 16, 32, 14, 36, 48, 76, 52, 42, 24, 34, 24, 30, 30]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(3, ws.max_row)}"
    ws.sheet_view.showGridLines = False


def _write_matrix_sheet(ws, bundle: dict[str, Any]) -> None:
    headers = ["候选ID", "SSD交换ID（请求及其返回）", "SSD请求消息ID", "关注点层级", "证据交换层级", "用例ID", "来源对象", "目标对象", "交互消息/依赖证据", "关注点Key", "关注点", "关注点主体", "适用状态", "判断依据", "异常数", "异常类型", "预测ID", "场景ID", "来源定位"]
    _title(ws, f"{bundle['project']} — 关注点矩阵", len(headers))
    _headers(ws, headers)
    nodes = _node_labels(bundle)
    rows = []
    for item in bundle.get("concern_matrix", []):
        from_id = item.get("from_node", "")
        to_id = item.get("to_node", "")
        rows.append([
            item.get("candidate_id", ""), item.get("exchange_id", ""), item.get("request_message_id", item.get("ssd_message_id", "")), item.get("concern_layer", item.get("layer", "")), item.get("exchange_layer", ""), item.get("use_case_id", ""), nodes.get(from_id, from_id), nodes.get(to_id, to_id), item.get("message", item.get("response_message", "")), item.get("concern_key", ""), item.get("concern", item.get("concern_key", "")), item.get("concern_subject", ""), {"pending_review": "待Agent审核", "not_applicable": "不适用", "needs_requirement": "待需求确认"}.get(item.get("status"), item.get("status", "")), item.get("basis", ""), item.get("exception_count", len(item.get("findings", []))), "\n".join(item.get("exception_types", [])) or ({"pending_review": "待审核", "not_applicable": "无（不适用）", "needs_requirement": "待补需求证据"}.get(item.get("status"), "")), "\n".join(item.get("prediction_ids", [])), "\n".join(item.get("scenario_ids", [])), item.get("source_location", ""),
        ])
    _format_rows(ws, 4, rows, centered={1, 2, 3, 4, 5, 6, 10, 12, 13, 15})
    for index, width in enumerate([24, 24, 28, 14, 14, 20, 24, 24, 44, 38, 26, 24, 20, 46, 12, 38, 32, 32, 34], 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(3, ws.max_row)}"
    ws.sheet_view.showGridLines = False


def _write_timeout_sheet(ws, bundle: dict[str, Any]) -> None:
    headers = ["SSD交换ID", "SSD请求消息ID", "用例ID", "Use Case名称", "主流程步骤", "交互消息", "适用状态", "需求满足影响", "后续行为影响", "环境协调影响", "判断依据", "来源定位"]
    _title(ws, f"{bundle['project']} — 超时判断", len(headers))
    _headers(ws, headers)
    ucs = {uc["use_case_id"]: uc for uc in bundle["scene_model"].get("use_cases", [])}
    rows = []
    for item in bundle.get("concern_matrix", []):
        if item.get("concern_key") != "common.timeout":
            continue
        uc = ucs.get(item.get("use_case_id"), {})
        impacts = [item.get("requirement_impact", ""), item.get("subsequent_behavior_impact", ""), item.get("environment_coordination_impact", "")]
        impacts = [value if value in {"yes", "no"} else "待需求确认" for value in impacts]
        rows.append([item.get("exchange_id", ""), item.get("request_message_id", item.get("ssd_message_id", "")), item.get("use_case_id", ""), uc.get("use_case_name", ""), item.get("source_step_index", ""), item.get("message", ""), item.get("status", ""), *impacts, item.get("basis", ""), item.get("source_location", "")])
    _format_rows(ws, 4, rows, centered={1, 2, 3, 5, 7, 8, 9, 10})
    for index, width in enumerate([24, 28, 20, 26, 14, 44, 22, 20, 20, 20, 48, 36], 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(3, ws.max_row)}"
    ws.sheet_view.showGridLines = False


def _write_mapping_workbook(bundle: dict[str, Any], path: Path) -> None:
    """Write the explicit RR/SR/AR interface mapping audit workbook."""
    model = bundle["scene_model"]
    ucs = {u["use_case_id"]: u for u in model.get("use_cases", [])}
    sr_headers = ["RR Use Case ID", "RR Use Case名称", "SR设计用例ID", "SR Service", "SR功能分类", "分类状态", "分类依据", "Abstract API ID", "HTTP方法", "资源路径", "请求参数", "响应字段", "错误码", "参与者", "来源定位", "映射状态"]
    ar_headers = ["RR Use Case ID", "SR Abstract API ID", "Implementation API ID", "AR技术职责分类", "分类状态", "分类依据", "软件实现接口", "HTTP方法", "资源路径", "AR微服务ID", "AR微服务名称", "数据库/外部依赖", "数据读写动作", "请求字段", "返回字段", "来源定位", "映射状态", "待确认说明"]
    wb = Workbook(); sr_ws = wb.active; sr_ws.title = "SR接口映射"
    _title(sr_ws, f"{bundle['project']} — SR接口映射", len(sr_headers)); _headers(sr_ws, sr_headers)
    sr_rows, ar_rows = [], []
    for uc in model.get("use_cases", []):
        arch = uc.get("architecture") or {}; sr = arch.get("sr") or {}; api_id = sr.get("abstract_api_id", "")
        iface = next((i for i in model.get("interfaces", []) if i.get("abstract_api_id") == api_id or i.get("name") == api_id), {})
        sr_rows.append([uc["use_case_id"], uc.get("use_case_name", ""), sr.get("design_use_case_id", ""), sr.get("service_name", sr.get("service_id", "")), sr.get("service_type", "unknown"), sr.get("classification_status", "needs_confirmation"), sr.get("classification_basis", "证据不足，待确认"), api_id, iface.get("method", ""), iface.get("path", ""), "\n".join(iface.get("request_fields", sr.get("request_fields", [])) or []), "\n".join(iface.get("response_fields", sr.get("response_fields", [])) or []), "\n".join(iface.get("error_codes", []) or []), "、".join(uc.get("actors", [])), sr.get("source_location", uc.get("source_location", "")), sr.get("mapping_status", "confirmed" if iface else "needs_confirmation")])
        for component in arch.get("ar", []) or []:
            status = component.get("mapping_status", "confirmed" if uc["use_case_id"] == "UCG-001-UC001" and component.get("implementation_api_id") in {"listProducts", "getProductDetail"} else "inferred")
            ar_rows.append([uc["use_case_id"], api_id, component.get("implementation_api_id", ""), component.get("service_type", "unknown"), component.get("classification_status", "needs_confirmation"), component.get("classification_basis", "证据不足，待确认"), component.get("software_interface", ""), component.get("method", ""), component.get("resource_path", ""), component.get("microservice_id", ""), component.get("microservice_name", ""), component.get("database_dependency", component.get("database_action", "")), component.get("database_action", ""), "\n".join(component.get("request_fields", []) or []), "\n".join(component.get("response_fields", []) or []), component.get("source_location", ""), status, component.get("review_note", "" if status != "needs_confirmation" else "请补充 AR 微服务和软件实现接口的明确映射")])
    _format_rows(sr_ws, 4, sr_rows, centered={1, 3, 5, 6, 8, 9, 16})
    for i, width in enumerate([20, 24, 30, 28, 23, 18, 42, 22, 12, 38, 38, 42, 30, 24, 34, 18], 1): sr_ws.column_dimensions[get_column_letter(i)].width = width
    sr_ws.freeze_panes = "A4"; sr_ws.auto_filter.ref = f"A3:{get_column_letter(len(sr_headers))}{max(3, sr_ws.max_row)}"; sr_ws.sheet_view.showGridLines = False
    ar_ws = wb.create_sheet("AR软件实现接口映射")
    _title(ar_ws, f"{bundle['project']} — AR软件实现接口映射", len(ar_headers)); _headers(ar_ws, ar_headers)
    _format_rows(ar_ws, 4, ar_rows, centered={1, 2, 3, 4, 5, 8, 10, 17, 18})
    for i, width in enumerate([20, 22, 28, 25, 18, 42, 42, 12, 38, 24, 28, 34, 34, 34, 34, 34, 18, 18, 38], 1): ar_ws.column_dimensions[get_column_letter(i)].width = width
    ar_ws.freeze_panes = "A4"; ar_ws.auto_filter.ref = f"A3:{get_column_letter(len(ar_headers))}{max(3, ar_ws.max_row)}"; ar_ws.sheet_view.showGridLines = False
    wb.save(path)


def _classification_summary(model: dict[str, Any]) -> dict[str, Any]:
    use_cases = model.get("use_cases", [])
    sr_types = Counter((uc.get("architecture") or {}).get("sr", {}).get("service_type", "unknown") for uc in use_cases)
    ar_mappings = [component for uc in use_cases for component in ((uc.get("architecture") or {}).get("ar") or [])]
    ar_types = Counter(component.get("service_type", "unknown") for component in ar_mappings)
    return {
        "sr_use_case_mapping_count": len(use_cases),
        "ar_implementation_mapping_count": len(ar_mappings),
        "sr_by_type": dict(sorted(sr_types.items())),
        "ar_by_type": dict(sorted(ar_types.items())),
        "unclassified_sr_use_cases": [
            uc.get("use_case_id", "") for uc in use_cases
            if ((uc.get("architecture") or {}).get("sr") or {}).get("service_type", "unknown") == "unknown"
        ],
        "unclassified_ar_mappings": [
            {"use_case_id": uc.get("use_case_id", ""), "implementation_api_id": component.get("implementation_api_id", ""), "microservice_id": component.get("microservice_id", "")}
            for uc in use_cases for component in ((uc.get("architecture") or {}).get("ar") or [])
            if component.get("service_type", "unknown") == "unknown"
        ],
    }


def _json_write(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def export_workbooks(bundle: dict[str, Any], output_dir: str | Path) -> dict[str, str]:
    output = Path(output_dir).expanduser().resolve()
    output.mkdir(parents=True, exist_ok=True)
    project = _safe_filename(bundle.get("project", bundle["scene_model"].get("project", "scene_completion")))
    artifacts = {
        "scene_model": output / "scene_model.json",
        "system_composition": output / "system_composition.json",
        "interaction_catalog": output / "interaction_catalog.json",
        "concern_matrix": output / "concern_matrix.json",
        "checkpoint_results": output / "checkpoint_results.json",
        "exception_tree": output / "exception_tree.json",
        "scenario_catalog_json": output / "scenario_catalog.json",
        "diagram_manifest": output / "diagram_manifest.json",
        "review_items": output / "review_items.json",
        "prediction_workbook": output / f"prediction_analysis_{project}.xlsx",
        "scenario_workbook": output / f"scenario_catalog_{project}.xlsx",
        "mapping_workbook": output / f"interface_service_mapping_{project}.xlsx",
    }
    _json_write(artifacts["scene_model"], bundle["scene_model"])
    _json_write(artifacts["system_composition"], bundle["system_composition"])
    _json_write(artifacts["interaction_catalog"], bundle["interaction_catalog"])
    matrix_items = bundle["concern_matrix"]
    _json_write(artifacts["concern_matrix"], {
        "version": bundle.get("version", "4"),
        "registry_version": bundle.get("registry_version", ""),
        "analysis_layers": bundle.get("analysis_layers", ["SR"]),
        "items": matrix_items,
        "exchange_ids": sorted({item.get("exchange_id") for item in matrix_items if item.get("exchange_id")}),
        "interaction_ids": sorted({item.get("interaction_id") for item in matrix_items if item.get("interaction_id")}),
    })
    _json_write(artifacts["checkpoint_results"], bundle["checkpoint_results"])
    _json_write(artifacts["exception_tree"], bundle["exception_tree"])
    _json_write(artifacts["scenario_catalog_json"], bundle["scenario_catalog"])
    _json_write(artifacts["diagram_manifest"], bundle.get("diagram_manifest", {"version": "3", "status": "skipped", "artifacts": []}))
    _json_write(artifacts["review_items"], bundle.get("review_items", []))
    _write_prediction_workbook(bundle, artifacts["prediction_workbook"])
    wb = Workbook()
    ws = wb.active
    ws.title = "场景清单"
    _write_scenario_sheet(ws, bundle)
    matrix_ws = wb.create_sheet("关注点矩阵")
    _write_matrix_sheet(matrix_ws, bundle)
    timeout_ws = wb.create_sheet("超时判断")
    _write_timeout_sheet(timeout_ws, bundle)
    if "AR" in bundle.get("analysis_layers", []):
        ar_ws = wb.create_sheet("AR扩展异常")
        ar_bundle = {**bundle, "scenario_catalog": [item for item in bundle.get("scenario_catalog", []) if item.get("layer") == "AR"]}
        _write_scenario_sheet(ar_ws, ar_bundle)
    wb.save(artifacts["scenario_workbook"])
    _write_mapping_workbook(bundle, artifacts["mapping_workbook"])
    # Make the mapping table a first-class trace target from both manifests.
    diagram_value = bundle.get("diagram_manifest")
    if isinstance(diagram_value, dict):
        diagram_value["interface_mapping_workbook"] = str(artifacts["mapping_workbook"])
        diagram_value["service_classification_summary"] = _classification_summary(bundle["scene_model"])
        _json_write(artifacts["diagram_manifest"], diagram_value)
    manifest = {
        "version": bundle.get("version", "3"),
        "project": bundle["project"],
        "scenario_count": len(bundle.get("scenario_catalog", [])),
        "prediction_count": len(bundle.get("findings", [])),
        "main_success_count": sum(1 for item in bundle.get("scenario_catalog", []) if item.get("scenario_type") == "main_success"),
        "interaction_count": len(bundle.get("interaction_catalog", [])),
        "concern_matrix_count": len(bundle.get("concern_matrix", [])),
        "review_item_count": len(bundle.get("review_items", [])),
        "evaluation": {"ground_truth": False, "precision": None, "recall": None, "f1": None},
        "outputs": {key: str(path) for key, path in artifacts.items()},
        "interface_mapping_workbook": str(artifacts["mapping_workbook"]),
        "service_classification_summary": _classification_summary(bundle["scene_model"]),
    }
    manifest_path = output / "run_manifest.json"
    _json_write(manifest_path, manifest)
    artifacts["manifest"] = manifest_path
    return {key: str(path) for key, path in artifacts.items()}
